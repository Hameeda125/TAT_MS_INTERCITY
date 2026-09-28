"""Stage 5: compact TAT computation for every paired workbook.

Run after phase_segmentation.py:

    python src/tat_computation.py

This version creates only TWO result sheets in every paired workbook:

1. 08_TAT_Phase_Comparison
   - EV and ICE total TAT
   - total distance
   - moving time and distance
   - average moving speed
   - stationary time
   - phase-wise EV/ICE time and distance in one table
   - reconciliation checks

2. 09_SpeedRange_Comparison
   - EV and ICE time spent in each speed range in one table

It does not crop, alter source rows, or perform phase segmentation.
"""

from __future__ import annotations

import math
import os
import re
import stat
import time
import warnings
from datetime import datetime
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config_loader import load_config


SCRIPT_VERSION = "2026-08-06_STAGE_5_MINIMAL_TWO_SHEETS_V1"

PAIR_INDEX_NAME = "pair_index.csv"
TAT_LOG_NAME = "TAT_Computation_Log.xlsx"

EV_SOURCE_SHEET = "01_EV_Cropped_Data"
ICE_SOURCE_SHEET = "02_ICE_Cropped_Data"

TAT_RESULT_SHEET = "08_TAT_Phase_Comparison"
SPEED_RESULT_SHEET = "09_SpeedRange_Comparison"

# Remove all older Stage 5 outputs on rerun.
LEGACY_STAGE5_SHEETS = [
    "08_EV_Phase_Time",
    "09_EV_Phase_Distance",
    "10_EV_SpeedRange_Time",
    "11_EV_TAT_Summary",
    "12_ICE_Phase_Time",
    "13_ICE_Phase_Distance",
    "14_ICE_SpeedRange_Time",
    "15_ICE_TAT_Summary",
    "16_EV_ICE_Phase_Time_Comparison",
    "17_EV_ICE_Phase_Distance_Comparison",
    "18_EV_ICE_TAT_Comparison",
    "19_EV_ICE_MovingSpeed_Comparison",
    "16_Phase_Time_Compare",
    "17_Phase_Distance_Compare",
    "17_Phase_Dist_Compare",
    "18_TAT_Compare",
    "19_MovingSpeed_Compare",
    "19_Moving_Speed_Compare",
    "12_EV_Speed_Range_Time",
    "13_ICE_Speed_Range_Time",
    "14_EV_ICE_TAT_Comparison",
]

TIMESTAMP = "Timestamp"
TIME_DIFFERENCE = "Time Difference (s)"
DISTANCE = "Haversine Distance (m)"
SPEED_RANGE = "Speed Range (km/h)"
PHASE = "phase"

EV_PHASES = [
    "CHARGING",
    "STOP",
    "IDLE",
    "INCH",
    "ACC",
    "DEC",
    "CRUISE",
]
ICE_PHASES = [
    "REFUELING",
    "STOP",
    "IDLE",
    "INCH",
    "ACC",
    "DEC",
    "CRUISE",
]
COMMON_PHASES = [
    "CHARGING / REFUELING",
    "STOP",
    "IDLE",
    "INCH",
    "ACC",
    "DEC",
    "CRUISE",
]
MOVING_PHASES = {"INCH", "ACC", "DEC", "CRUISE"}

HEADER_BLUE = "1F4E78"
SECTION_BLUE = "D9EAF7"
TOTAL_GREEN = "E2F0D9"
WHITE = "FFFFFF"
THIN_GREY = Side(style="thin", color="B7B7B7")
BORDER = Border(
    left=THIN_GREY,
    right=THIN_GREY,
    top=THIN_GREY,
    bottom=THIN_GREY,
)

warnings.filterwarnings(
    "ignore",
    message=r"Title is more than 31 characters.*",
    category=UserWarning,
    module=r"openpyxl\.workbook\.child",
)


def _clean_text(value: Any) -> str:
    if value is None or value is pd.NA:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def _normalise_phase(value: Any) -> str:
    text = _clean_text(value).upper().replace(" ", "_")
    aliases = {
        "INCHING": "INCH",
        "ACCELERATION": "ACC",
        "DECELERATION": "DEC",
        "CRUISING": "CRUISE",
        "CHARGE": "CHARGING",
        "REFUELLING": "REFUELING",
        "REFUEL": "REFUELING",
    }
    return aliases.get(text, text)


def _read_csv(path: Path) -> pd.DataFrame:
    last_error = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return pd.read_csv(path, encoding=encoding)
        except UnicodeDecodeError as error:
            last_error = error
    raise ValueError(f"Could not decode {path}: {last_error}")


def _parse_datetime(series: pd.Series) -> pd.Series:
    try:
        parsed = pd.to_datetime(
            series,
            errors="coerce",
            dayfirst=True,
            format="mixed",
        )
    except TypeError:
        parsed = pd.to_datetime(
            series,
            errors="coerce",
            dayfirst=True,
        )

    unresolved = (
        parsed.isna()
        & series.notna()
        & series.astype(str).str.strip().ne("")
    )
    if bool(unresolved.any()):
        try:
            parsed.loc[unresolved] = pd.to_datetime(
                series.loc[unresolved],
                errors="coerce",
                dayfirst=False,
                format="mixed",
            )
        except TypeError:
            parsed.loc[unresolved] = pd.to_datetime(
                series.loc[unresolved],
                errors="coerce",
                dayfirst=False,
            )
    return parsed


def _parse_number(series: pd.Series) -> pd.Series:
    if pd.api.types.is_numeric_dtype(series.dtype):
        return pd.to_numeric(series, errors="coerce")

    text = series.astype("string").str.strip()
    text = text.str.replace("\u2212", "-", regex=False)
    text = text.str.replace(",", "", regex=False)
    extracted = text.str.extract(
        r"([-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?)",
        expand=False,
    )
    return pd.to_numeric(extracted, errors="coerce")


def _resolve_workbook_path(
    row: pd.Series,
    results_root: Path,
) -> Path:
    configured = _clean_text(row.get("pair_workbook_path"))
    if configured:
        candidate = Path(configured)
        if candidate.exists() and candidate.is_file():
            return candidate

    workbook_name = _clean_text(row.get("pair_workbook_name"))
    date_folder = _clean_text(row.get("output_date_folder"))
    if workbook_name and date_folder:
        candidate = results_root / "Paired" / date_folder / workbook_name
        if candidate.exists() and candidate.is_file():
            return candidate

    pair_id = _clean_text(row.get("pair_id"))
    matches = [
        path
        for path in (results_root / "Paired").rglob(f"{pair_id}_*.xlsx")
        if path.is_file()
        and not path.name.startswith(("~$", "."))
    ]
    if matches:
        return max(matches, key=lambda path: path.stat().st_mtime)

    raise FileNotFoundError(
        f"Pair workbook not found for {pair_id or workbook_name}"
    )


def _prepare_vehicle_data(
    source: pd.DataFrame,
    vehicle_type: str,
) -> pd.DataFrame:
    required = [
        TIMESTAMP,
        TIME_DIFFERENCE,
        DISTANCE,
        SPEED_RANGE,
        PHASE,
    ]
    missing = [column for column in required if column not in source.columns]
    if missing:
        raise ValueError(
            f"{vehicle_type} sheet is missing: " + ", ".join(missing)
        )

    frame = source[required].copy()
    frame[TIMESTAMP] = _parse_datetime(frame[TIMESTAMP])
    frame[TIME_DIFFERENCE] = _parse_number(frame[TIME_DIFFERENCE])
    frame[DISTANCE] = _parse_number(frame[DISTANCE])
    frame[SPEED_RANGE] = (
        frame[SPEED_RANGE]
        .astype("string")
        .fillna("")
        .str.strip()
    )
    frame.loc[frame[SPEED_RANGE].eq(""), SPEED_RANGE] = "UNSPECIFIED"
    frame[PHASE] = frame[PHASE].map(_normalise_phase)

    if frame.empty:
        raise ValueError(f"{vehicle_type} sheet is empty")
    if frame[TIMESTAMP].isna().any():
        raise ValueError(
            f"{vehicle_type} sheet contains invalid Timestamp values"
        )
    if frame[[TIME_DIFFERENCE, DISTANCE]].isna().any(axis=None):
        raise ValueError(
            f"{vehicle_type} sheet contains invalid time or distance values"
        )
    if (frame[TIME_DIFFERENCE] < 0).any():
        raise ValueError(f"{vehicle_type} sheet contains negative time")
    if (frame[DISTANCE] < 0).any():
        raise ValueError(f"{vehicle_type} sheet contains negative distance")

    positive = frame[TIME_DIFFERENCE] > 0
    if frame.loc[positive, PHASE].eq("").any():
        raise ValueError(
            f"{vehicle_type} has positive-duration rows without phase"
        )

    allowed = set(EV_PHASES if vehicle_type == "EV" else ICE_PHASES)
    unexpected = sorted(set(frame.loc[positive, PHASE]) - allowed)
    if unexpected:
        raise ValueError(
            f"{vehicle_type} has unexpected phase(s): "
            + ", ".join(unexpected)
        )

    return frame.reset_index(drop=True)


def _speed_range_sort_key(value: Any) -> tuple[float, float, str]:
    text = _clean_text(value)
    if text.startswith(">"):
        numbers = re.findall(r"\d+(?:\.\d+)?", text)
        lower = float(numbers[0]) if numbers else math.inf
        return (lower, math.inf, text)

    numbers = [
        float(number)
        for number in re.findall(r"\d+(?:\.\d+)?", text)
    ]
    if len(numbers) >= 2:
        return (numbers[0], numbers[1], text)
    if len(numbers) == 1:
        return (numbers[0], numbers[0], text)
    return (math.inf, math.inf, text)


def _compute_metrics(
    frame: pd.DataFrame,
    vehicle_type: str,
) -> dict:
    phase_order = EV_PHASES if vehicle_type == "EV" else ICE_PHASES
    positive = frame[TIME_DIFFERENCE] > 0

    total_time_s = float(frame.loc[positive, TIME_DIFFERENCE].sum())
    total_distance_m = float(frame.loc[positive, DISTANCE].sum())

    trip_start = frame[TIMESTAMP].iloc[0]
    trip_end = frame[TIMESTAMP].iloc[-1]
    elapsed_s = float((trip_end - trip_start).total_seconds())

    grouped = (
        frame.loc[positive]
        .groupby(PHASE)
        .agg(
            time_s=(TIME_DIFFERENCE, "sum"),
            distance_m=(DISTANCE, "sum"),
            intervals=(PHASE, "size"),
        )
    )

    phase_values: dict[str, dict] = {}
    for phase_name in phase_order:
        if phase_name in grouped.index:
            time_s = float(grouped.at[phase_name, "time_s"])
            distance_m = float(grouped.at[phase_name, "distance_m"])
            intervals = int(grouped.at[phase_name, "intervals"])
        else:
            time_s = 0.0
            distance_m = 0.0
            intervals = 0

        phase_values[phase_name] = {
            "time_s": time_s,
            "time_min": time_s / 60.0,
            "distance_m": distance_m,
            "distance_km": distance_m / 1000.0,
            "intervals": intervals,
            "time_pct": (
                time_s / total_time_s * 100.0
                if total_time_s > 0
                else 0.0
            ),
            "distance_pct": (
                distance_m / total_distance_m * 100.0
                if total_distance_m > 0
                else 0.0
            ),
        }

    moving_mask = positive & frame[PHASE].isin(MOVING_PHASES)
    moving_time_s = float(
        frame.loc[moving_mask, TIME_DIFFERENCE].sum()
    )
    moving_distance_m = float(
        frame.loc[moving_mask, DISTANCE].sum()
    )

    average_moving_speed_kmh = (
        moving_distance_m / moving_time_s * 3.6
        if moving_time_s > 0
        else 0.0
    )

    overall_average_speed_kmh = (
        total_distance_m / total_time_s * 3.6
        if total_time_s > 0
        else 0.0
    )

    speed_grouped = (
        frame.loc[positive]
        .groupby(SPEED_RANGE)[TIME_DIFFERENCE]
        .sum()
        .to_dict()
    )

    phase_time_total_s = sum(
        phase_values[phase]["time_s"]
        for phase in phase_order
    )
    phase_distance_total_m = sum(
        phase_values[phase]["distance_m"]
        for phase in phase_order
    )

    return {
        "vehicle_type": vehicle_type,
        "trip_start": trip_start,
        "trip_end": trip_end,
        "total_time_s": total_time_s,
        "total_time_min": total_time_s / 60.0,
        "total_time_hr": total_time_s / 3600.0,
        "total_distance_m": total_distance_m,
        "total_distance_km": total_distance_m / 1000.0,
        "moving_time_s": moving_time_s,
        "moving_time_min": moving_time_s / 60.0,
        "moving_distance_m": moving_distance_m,
        "moving_distance_km": moving_distance_m / 1000.0,
        "stationary_time_s": total_time_s - moving_time_s,
        "stationary_time_min": (total_time_s - moving_time_s) / 60.0,
        "average_moving_speed_kmh": average_moving_speed_kmh,
        "overall_average_speed_kmh": overall_average_speed_kmh,
        "phase_values": phase_values,
        "speed_grouped": speed_grouped,
        "input_time_error_s": total_time_s - elapsed_s,
        "phase_time_error_s": phase_time_total_s - total_time_s,
        "phase_distance_error_m": (
            phase_distance_total_m - total_distance_m
        ),
    }


def _phase_comparison_table(
    ev: dict,
    ice: dict,
) -> pd.DataFrame:
    rows = []

    for display_phase in COMMON_PHASES:
        if display_phase == "CHARGING / REFUELING":
            ev_phase = "CHARGING"
            ice_phase = "REFUELING"
        else:
            ev_phase = display_phase
            ice_phase = display_phase

        ev_values = ev["phase_values"][ev_phase]
        ice_values = ice["phase_values"][ice_phase]

        rows.append(
            {
                "Phase": display_phase,
                "EV Time (min)": ev_values["time_min"],
                "ICE Time (min)": ice_values["time_min"],
                "Time Difference EV-ICE (min)": (
                    ev_values["time_min"]
                    - ice_values["time_min"]
                ),
                "EV Distance (km)": ev_values["distance_km"],
                "ICE Distance (km)": ice_values["distance_km"],
                "Distance Difference EV-ICE (km)": (
                    ev_values["distance_km"]
                    - ice_values["distance_km"]
                ),
                "EV TAT Contribution (%)": ev_values["time_pct"],
                "ICE TAT Contribution (%)": ice_values["time_pct"],
                "EV Distance Contribution (%)": ev_values["distance_pct"],
                "ICE Distance Contribution (%)": ice_values["distance_pct"],
            }
        )

    return pd.DataFrame(rows)


def _overall_comparison_table(
    ev: dict,
    ice: dict,
) -> pd.DataFrame:
    metrics = [
        (
            "Total TAT (min)",
            ev["total_time_min"],
            ice["total_time_min"],
        ),
        (
            "Total TAT (hr)",
            ev["total_time_hr"],
            ice["total_time_hr"],
        ),
        (
            "Total Distance (km)",
            ev["total_distance_km"],
            ice["total_distance_km"],
        ),
        (
            "Moving Time (min)",
            ev["moving_time_min"],
            ice["moving_time_min"],
        ),
        (
            "Moving Distance (km)",
            ev["moving_distance_km"],
            ice["moving_distance_km"],
        ),
        (
            "Average Moving Speed (km/h)",
            ev["average_moving_speed_kmh"],
            ice["average_moving_speed_kmh"],
        ),
        (
            "Overall Average Speed (km/h)",
            ev["overall_average_speed_kmh"],
            ice["overall_average_speed_kmh"],
        ),
        (
            "Stationary Time (min)",
            ev["stationary_time_min"],
            ice["stationary_time_min"],
        ),
        (
            "Input Time Error (s)",
            ev["input_time_error_s"],
            ice["input_time_error_s"],
        ),
        (
            "Phase Time Error (s)",
            ev["phase_time_error_s"],
            ice["phase_time_error_s"],
        ),
        (
            "Phase Distance Error (m)",
            ev["phase_distance_error_m"],
            ice["phase_distance_error_m"],
        ),
    ]

    rows = []
    for metric, ev_value, ice_value in metrics:
        rows.append(
            {
                "Metric": metric,
                "EV": ev_value,
                "ICE": ice_value,
                "EV - ICE": ev_value - ice_value,
                "EV vs ICE (%)": (
                    (ev_value - ice_value) / ice_value * 100.0
                    if abs(ice_value) > 1e-12
                    else np.nan
                ),
            }
        )

    return pd.DataFrame(rows)


def _speed_comparison_table(
    ev: dict,
    ice: dict,
) -> pd.DataFrame:
    ranges = sorted(
        set(ev["speed_grouped"]) | set(ice["speed_grouped"]),
        key=_speed_range_sort_key,
    )

    ev_total = sum(ev["speed_grouped"].values())
    ice_total = sum(ice["speed_grouped"].values())

    rows = []
    for speed_range in ranges:
        ev_s = float(ev["speed_grouped"].get(speed_range, 0.0))
        ice_s = float(ice["speed_grouped"].get(speed_range, 0.0))

        rows.append(
            {
                "Speed Range": speed_range,
                "EV Time (s)": ev_s,
                "ICE Time (s)": ice_s,
                "EV Time (min)": ev_s / 60.0,
                "ICE Time (min)": ice_s / 60.0,
                "Time Difference EV-ICE (min)": (
                    ev_s - ice_s
                ) / 60.0,
                "EV Time Contribution (%)": (
                    ev_s / ev_total * 100.0
                    if ev_total > 0
                    else 0.0
                ),
                "ICE Time Contribution (%)": (
                    ice_s / ice_total * 100.0
                    if ice_total > 0
                    else 0.0
                ),
            }
        )

    rows.append(
        {
            "Speed Range": "TOTAL",
            "EV Time (s)": ev_total,
            "ICE Time (s)": ice_total,
            "EV Time (min)": ev_total / 60.0,
            "ICE Time (min)": ice_total / 60.0,
            "Time Difference EV-ICE (min)": (
                ev_total - ice_total
            ) / 60.0,
            "EV Time Contribution (%)": 100.0,
            "ICE Time Contribution (%)": 100.0,
        }
    )

    return pd.DataFrame(rows)


def _write_section_title(
    worksheet,
    row_number: int,
    title: str,
    max_column: int,
) -> None:
    worksheet.merge_cells(
        start_row=row_number,
        start_column=1,
        end_row=row_number,
        end_column=max_column,
    )
    cell = worksheet.cell(row_number, 1, title)
    cell.fill = PatternFill(
        fill_type="solid",
        fgColor=HEADER_BLUE,
    )
    cell.font = Font(
        color=WHITE,
        bold=True,
        size=12,
    )
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )
    worksheet.row_dimensions[row_number].height = 24


def _write_table(
    worksheet,
    start_row: int,
    dataframe: pd.DataFrame,
    total_row: bool = False,
) -> int:
    header_row = start_row

    for column_number, column_name in enumerate(
        dataframe.columns,
        start=1,
    ):
        cell = worksheet.cell(
            header_row,
            column_number,
            column_name,
        )
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor=SECTION_BLUE,
        )
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = BORDER

    for row_offset, values in enumerate(
        dataframe.itertuples(index=False, name=None),
        start=1,
    ):
        excel_row = header_row + row_offset
        for column_number, value in enumerate(
            values,
            start=1,
        ):
            if isinstance(value, np.generic):
                value = value.item()
            elif value is pd.NA:
                value = None
            elif isinstance(value, float) and math.isnan(value):
                value = None

            cell = worksheet.cell(
                excel_row,
                column_number,
                value,
            )
            cell.border = BORDER
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

            if isinstance(value, (int, float)) and not isinstance(value, bool):
                cell.number_format = "0.000"

    last_row = header_row + len(dataframe)

    if total_row and len(dataframe) > 0:
        for cell in worksheet[last_row]:
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor=TOTAL_GREEN,
            )
            cell.font = Font(bold=True)

    return last_row


def _format_result_sheet(worksheet) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False

    for column_number in range(1, worksheet.max_column + 1):
        longest = 12
        for row_number in range(1, min(worksheet.max_row, 100) + 1):
            longest = max(
                longest,
                len(
                    _clean_text(
                        worksheet.cell(
                            row_number,
                            column_number,
                        ).value
                    )
                ),
            )
        worksheet.column_dimensions[
            get_column_letter(column_number)
        ].width = min(max(longest + 2, 14), 32)


def _write_tat_result_sheet(
    workbook,
    ev_metrics: dict,
    ice_metrics: dict,
) -> None:
    if TAT_RESULT_SHEET in workbook.sheetnames:
        workbook.remove(workbook[TAT_RESULT_SHEET])

    worksheet = workbook.create_sheet(TAT_RESULT_SHEET)

    overall = _overall_comparison_table(
        ev_metrics,
        ice_metrics,
    )
    phase = _phase_comparison_table(
        ev_metrics,
        ice_metrics,
    )

    _write_section_title(
        worksheet,
        1,
        "EV–ICE Overall TAT and Operational Comparison",
        len(overall.columns),
    )
    overall_end = _write_table(
        worksheet,
        2,
        overall,
        total_row=False,
    )

    phase_title_row = overall_end + 3
    _write_section_title(
        worksheet,
        phase_title_row,
        "EV–ICE Phase Time, Distance and Contribution Comparison",
        len(phase.columns),
    )
    _write_table(
        worksheet,
        phase_title_row + 1,
        phase,
        total_row=False,
    )

    _format_result_sheet(worksheet)


def _write_speed_result_sheet(
    workbook,
    ev_metrics: dict,
    ice_metrics: dict,
) -> None:
    if SPEED_RESULT_SHEET in workbook.sheetnames:
        workbook.remove(workbook[SPEED_RESULT_SHEET])

    worksheet = workbook.create_sheet(SPEED_RESULT_SHEET)
    speed = _speed_comparison_table(
        ev_metrics,
        ice_metrics,
    )

    _write_section_title(
        worksheet,
        1,
        "EV–ICE Speed Range and Time-Spent Comparison",
        len(speed.columns),
    )
    _write_table(
        worksheet,
        2,
        speed,
        total_row=True,
    )

    _format_result_sheet(worksheet)


def _save_workbook(
    path: Path,
    ev_metrics: dict,
    ice_metrics: dict,
) -> None:
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message=r"Title is more than 31 characters.*",
            category=UserWarning,
        )
        workbook = load_workbook(
            path,
            data_only=False,
        )

    temporary_path = None

    try:
        for sheet_name in LEGACY_STAGE5_SHEETS:
            if sheet_name in workbook.sheetnames:
                workbook.remove(workbook[sheet_name])

        _write_tat_result_sheet(
            workbook,
            ev_metrics,
            ice_metrics,
        )
        _write_speed_result_sheet(
            workbook,
            ev_metrics,
            ice_metrics,
        )

        with NamedTemporaryFile(
            prefix=f".{path.stem}_tat_",
            suffix=".xlsx",
            dir=path.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)

        workbook.save(temporary_path)
        workbook.close()
        workbook = None

        try:
            os.chmod(
                path,
                stat.S_IREAD | stat.S_IWRITE,
            )
        except OSError:
            pass

        os.replace(temporary_path, path)

    finally:
        if workbook is not None:
            workbook.close()

        if (
            temporary_path is not None
            and temporary_path.exists()
        ):
            temporary_path.unlink(missing_ok=True)


def _atomic_write_csv(
    path: Path,
    dataframe: pd.DataFrame,
) -> None:
    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with NamedTemporaryFile(
        prefix=f".{path.stem}_",
        suffix=path.suffix,
        dir=path.parent,
        delete=False,
    ) as temporary_file:
        temporary_path = Path(temporary_file.name)

    try:
        dataframe.to_csv(
            temporary_path,
            index=False,
            encoding="utf-8-sig",
        )
        os.replace(temporary_path, path)
    except Exception:
        temporary_path.unlink(missing_ok=True)
        raise


def _format_duration(seconds: float) -> str:
    total_seconds = max(0, int(round(seconds)))
    hours, remainder = divmod(total_seconds, 3600)
    minutes, seconds = divmod(remainder, 60)
    return (
        f"{hours:02d}:"
        f"{minutes:02d}:"
        f"{seconds:02d}"
    )


def run_tat_computation(config: dict) -> dict:
    results_root = Path(
        config["project"]["results_folder"]
    )
    system_root = results_root / "System"
    pair_index_path = (
        system_root / PAIR_INDEX_NAME
    )

    if not pair_index_path.exists():
        raise FileNotFoundError(
            f"Pair index not found: {pair_index_path}"
        )

    pair_index = _read_csv(pair_index_path)
    if pair_index.empty:
        raise ValueError("pair_index.csv is empty")

    output_columns = [
        "tat_computation_version",
        "tat_computation_status",
        "tat_computation_completed_at",
        "tat_computation_error",
        "ev_total_tat_seconds",
        "ice_total_tat_seconds",
        "ev_total_distance_km",
        "ice_total_distance_km",
        "ev_moving_time_seconds",
        "ice_moving_time_seconds",
        "ev_average_moving_speed_kmh",
        "ice_average_moving_speed_kmh",
    ]
    for column in output_columns:
        if column not in pair_index.columns:
            pair_index[column] = pd.NA

    print("=" * 80)
    print("STAGE 5: COMPACT TAT COMPUTATION")
    print("=" * 80)
    print(f"Script version: {SCRIPT_VERSION}")
    print(f"Pair workbooks: {len(pair_index)}")
    print("Result sheets per workbook: 2")
    print(f"  - {TAT_RESULT_SHEET}")
    print(f"  - {SPEED_RESULT_SHEET}")
    print("Moving phases: INCH + ACC + DEC + CRUISE")
    print()

    started = time.monotonic()
    audit_rows = []

    for sequence, (index, row) in enumerate(
        pair_index.iterrows(),
        start=1,
    ):
        pair_started = time.monotonic()
        pair_id = _clean_text(row.get("pair_id"))
        workbook_path = _resolve_workbook_path(
            row,
            results_root,
        )

        print(
            f"[{sequence}/{len(pair_index)}] "
            f"{pair_id} | {workbook_path.name}",
            flush=True,
        )

        status = "COMPLETED"
        error_text = "NONE"

        try:
            with pd.ExcelFile(
                workbook_path,
                engine="openpyxl",
            ) as excel:
                ev_source = pd.read_excel(
                    excel,
                    sheet_name=EV_SOURCE_SHEET,
                )
                ice_source = pd.read_excel(
                    excel,
                    sheet_name=ICE_SOURCE_SHEET,
                )

            ev_frame = _prepare_vehicle_data(
                ev_source,
                "EV",
            )
            ice_frame = _prepare_vehicle_data(
                ice_source,
                "ICE",
            )

            ev_metrics = _compute_metrics(
                ev_frame,
                "EV",
            )
            ice_metrics = _compute_metrics(
                ice_frame,
                "ICE",
            )

            _save_workbook(
                workbook_path,
                ev_metrics,
                ice_metrics,
            )

            completed_at = datetime.now()

            assignments = {
                "tat_computation_version": SCRIPT_VERSION,
                "tat_computation_status": status,
                "tat_computation_completed_at": (
                    completed_at.strftime(
                        "%d-%m-%Y %H:%M:%S"
                    )
                ),
                "tat_computation_error": "NONE",
                "ev_total_tat_seconds": ev_metrics[
                    "total_time_s"
                ],
                "ice_total_tat_seconds": ice_metrics[
                    "total_time_s"
                ],
                "ev_total_distance_km": ev_metrics[
                    "total_distance_km"
                ],
                "ice_total_distance_km": ice_metrics[
                    "total_distance_km"
                ],
                "ev_moving_time_seconds": ev_metrics[
                    "moving_time_s"
                ],
                "ice_moving_time_seconds": ice_metrics[
                    "moving_time_s"
                ],
                "ev_average_moving_speed_kmh": ev_metrics[
                    "average_moving_speed_kmh"
                ],
                "ice_average_moving_speed_kmh": ice_metrics[
                    "average_moving_speed_kmh"
                ],
            }

            for column, value in assignments.items():
                pair_index.at[index, column] = value

            audit_rows.append(
                {
                    "pair_id": pair_id,
                    "pair_workbook_path": str(workbook_path),
                    "status": status,
                    "ev_total_tat_min": ev_metrics[
                        "total_time_min"
                    ],
                    "ice_total_tat_min": ice_metrics[
                        "total_time_min"
                    ],
                    "ev_total_distance_km": ev_metrics[
                        "total_distance_km"
                    ],
                    "ice_total_distance_km": ice_metrics[
                        "total_distance_km"
                    ],
                    "ev_average_moving_speed_kmh": ev_metrics[
                        "average_moving_speed_kmh"
                    ],
                    "ice_average_moving_speed_kmh": ice_metrics[
                        "average_moving_speed_kmh"
                    ],
                    "error": "NONE",
                }
            )

        except Exception as error:
            status = "ERROR"
            error_text = str(error)

            pair_index.at[
                index,
                "tat_computation_version",
            ] = SCRIPT_VERSION
            pair_index.at[
                index,
                "tat_computation_status",
            ] = status
            pair_index.at[
                index,
                "tat_computation_completed_at",
            ] = datetime.now().strftime(
                "%d-%m-%Y %H:%M:%S"
            )
            pair_index.at[
                index,
                "tat_computation_error",
            ] = error_text

            audit_rows.append(
                {
                    "pair_id": pair_id,
                    "pair_workbook_path": str(workbook_path),
                    "status": status,
                    "error": error_text,
                }
            )

        elapsed = time.monotonic() - started
        eta = (
            elapsed
            / sequence
            * (len(pair_index) - sequence)
        )

        print(
            f"    Finished | Status: {status} | "
            f"This pair: "
            f"{_format_duration(time.monotonic() - pair_started)} | "
            f"Elapsed: {_format_duration(elapsed)} | "
            f"ETA: {_format_duration(eta)}",
            flush=True,
        )

        if status == "ERROR":
            print(
                f"    Error: {error_text}",
                flush=True,
            )

        print("-" * 80, flush=True)

    _atomic_write_csv(
        pair_index_path,
        pair_index,
    )

    log_path = system_root / TAT_LOG_NAME
    with pd.ExcelWriter(
        log_path,
        engine="openpyxl",
    ) as writer:
        pd.DataFrame(
            [
                {
                    "Metric": "Script version",
                    "Value": SCRIPT_VERSION,
                },
                {
                    "Metric": "Pairs received",
                    "Value": len(pair_index),
                },
                {
                    "Metric": "Pairs completed",
                    "Value": int(
                        pair_index[
                            "tat_computation_status"
                        ]
                        .astype(str)
                        .eq("COMPLETED")
                        .sum()
                    ),
                },
                {
                    "Metric": "Pairs with errors",
                    "Value": int(
                        pair_index[
                            "tat_computation_status"
                        ]
                        .astype(str)
                        .eq("ERROR")
                        .sum()
                    ),
                },
                {
                    "Metric": "Result sheets per workbook",
                    "Value": 2,
                },
                {
                    "Metric": "Moving phases",
                    "Value": "INCH + ACC + DEC + CRUISE",
                },
                {
                    "Metric": "Average moving speed formula",
                    "Value": (
                        "Moving distance / Moving time × 3.6"
                    ),
                },
            ]
        ).to_excel(
            writer,
            sheet_name="Run Summary",
            index=False,
        )

        pd.DataFrame(audit_rows).to_excel(
            writer,
            sheet_name="Pair TAT Audit",
            index=False,
        )

    completed_count = int(
        pair_index[
            "tat_computation_status"
        ]
        .astype(str)
        .eq("COMPLETED")
        .sum()
    )
    error_count = int(
        pair_index[
            "tat_computation_status"
        ]
        .astype(str)
        .eq("ERROR")
        .sum()
    )

    print()
    print("=" * 80)
    print("STAGE 5 PROCESSING SUMMARY")
    print("=" * 80)
    print(
        f"Completed pairs: "
        f"{completed_count}/{len(pair_index)}"
    )
    print(f"Pairs with errors: {error_count}")
    print(
        f"Total elapsed: "
        f"{_format_duration(time.monotonic() - started)}"
    )
    print(f"TAT log: {log_path}")
    print("=" * 80)

    if error_count:
        raise RuntimeError(
            f"Stage 5 completed with {error_count} error(s). "
            f"Review {log_path}."
        )

    return {
        "pair_count": len(pair_index),
        "completed_count": completed_count,
        "error_count": error_count,
        "pair_index_path": str(pair_index_path),
        "tat_log_path": str(log_path),
    }


def main() -> None:
    try:
        result = run_tat_computation(
            load_config()
        )

        print(
            "\nTAT STAGE 5 COMPLETED SUCCESSFULLY"
        )
        print(
            f"Pair workbooks processed: "
            f"{result['completed_count']}"
        )
        print(
            f"Updated pair index: "
            f"{result['pair_index_path']}"
        )
        print(
            f"TAT log: "
            f"{result['tat_log_path']}"
        )

    except Exception as error:
        print(
            f"\nTAT STAGE 5 FAILED | {error}"
        )
        raise SystemExit(1)


if __name__ == "__main__":
    main()