"""Reconcile one telemetry workbook with one intended trip-log row.

This stage is deliberately conservative.  It performs a global one-to-one
assignment between the permanent telemetry index and the detailed trip-log
sheet, then verifies each matched trip boundary from the full cleaned GPS
series.  Time differences are evidence, not a hard matching gate: route,
direction, date, vehicle type, timestamps, movement blocks, and any available
logged distance are considered together.

The module never overwrites raw or cleaned telemetry.  It writes a separate
``verified_trip_index.csv`` and ``Trip_Log_Reconciliation.xlsx``.  Strict and
non-strict EV-ICE labels belong to the later pairing stage; reconciliation
statuses only describe whether one intended trip was identified safely.
"""

from __future__ import annotations

import hashlib
import os
import re
import unicodedata
from datetime import date, datetime, time, timedelta
from difflib import SequenceMatcher
from pathlib import Path
from tempfile import NamedTemporaryFile

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config_loader import load_config
from trip_metadata import (
    calculate_trip_distance_km,
    haversine_distance_km,
    load_cleaned_telemetry,
)


DEFAULT_TRIP_LOG_PATH = Path(
    r"D:\TAT_Project\TAT Data Collection Intercity.xlsx"
)
TRIP_LOG_SHEET_NAME = "Sheet1"
VERIFIED_INDEX_FILE_NAME = "verified_trip_index.csv"
AUDIT_FILE_NAME = "Trip_Log_Reconciliation.xlsx"
SCRIPT_VERSION = "2026-09-29_RECONCILIATION_V3_INDEX_FRESHNESS_PACKET_TIME"

REQUIRED_INDEX_COLUMNS = {
    "trip_id",
    "file_name",
    "relative_path",
    "cleaned_file_path",
    "vehicle_type",
    "start_time",
    "end_time",
}

REQUIRED_LOG_COLUMNS = {
    "Date",
    "Vehicle Type",
    "Start Location",
    "End Location",
    "Start time",
    "End Time",
}

PLACEHOLDER_TEXT = {
    "",
    "-",
    "--",
    "—",
    "na",
    "n/a",
    "nan",
    "nil",
    "none",
    "not available",
    "20",
}

# Limited, auditable spelling aliases observed in the supplied files/log.
# KCBT and CMBT intentionally remain different termini.
LOCATION_ALIASES = {
    "cmbt": "cmbt",
    "cbmt": "cmbt",
    "cbmtbusstand": "cmbt",
    "cmbtbusstand": "cmbt",
    "koyambedu": "cmbt",
    "koyambedubusstand": "cmbt",
    "kcbt": "kcbt",
    "kilambakkam": "kcbt",
    "sholinganallur": "sholinganallur",
    "shollinganallur": "sholinganallur",
    "sholliganallur": "sholinganallur",
    "silkboard": "silkboard",
    "maduravoyal": "maduravoyal",
    "madhuravoyal": "maduravoyal",
    "maduravayol": "maduravoyal",
    "tnagar": "tnagar",
    "thiruporur": "thiruporur",
    "thirupporur": "thiruporur",
    "thirupur": "thiruporur",
    "tiruporur": "thiruporur",
    "pondy": "pondicherry",
    "pondi": "pondicherry",
    "pondicherry": "pondicherry",
    "puducherry": "pondicherry",
    "kia": "kia",
    "kempegowdainternationalairport": "kia",
    "kadugodibusstand": "kadugodi",
    "kadugodi": "kadugodi",
    "whitefield": "whitefield",
    "hebbal": "hebbal",
    "bellandur": "bellandur",
    "avrroundanasalem": "avrroundanasalem",
    "avrroundana": "avrroundanasalem",
    "salem": "avrroundanasalem",
    "chennai": "chennai",
    "bangalore": "bangalore",
    "bengaluru": "bangalore",
    "banglore": "bangalore",
    "madiwala": "madiwala",
    "lalbag": "lalbag",
    "lalbagh": "lalbag",
}

OPTIONAL_DISTANCE_COLUMNS = [
    "Distance (km)",
    "Distance(km)",
    "Distance",
    "Trip Distance (km)",
    "Trip Distance",
    "Distance Travelled (km)",
    "Distance Travelled",
]


def _clean_text(value):
    """Return stripped text, or an empty string for placeholders."""

    if value is None or pd.isna(value):
        return ""

    text_value = str(value).strip()
    if text_value.casefold() in PLACEHOLDER_TEXT:
        return ""

    return text_value


def _ascii_compact(value):
    """Create a case-insensitive alphanumeric comparison key."""

    text_value = _clean_text(value)
    normalized = unicodedata.normalize("NFKD", text_value)
    ascii_text = "".join(
        character
        for character in normalized
        if not unicodedata.combining(character)
    )
    return re.sub(r"[^a-z0-9]+", "", ascii_text.casefold())


def canonical_location(value):
    """Normalize one written location without online geocoding."""

    compact = _ascii_compact(value)
    if not compact:
        return ""
    return LOCATION_ALIASES.get(compact, compact)


def _vehicle_type(value):
    """Map log/index labels to EV, ICE, or UNKNOWN."""

    compact = _ascii_compact(value)
    if compact in {"ev", "electric", "electricvehicle"}:
        return "EV"
    if compact in {"ice", "diesel", "combustion", "internalcombustion"}:
        return "ICE"
    return "UNKNOWN"


def _parse_log_date(value):
    """Parse a real Excel date, serial date, or day-first date string."""

    if value is None or pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.date()
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        numeric = float(value)
        if numeric >= 20_000:
            parsed = pd.Timestamp("1899-12-30") + pd.to_timedelta(
                numeric, unit="D"
            )
            return parsed.date()

    text_value = _clean_text(value)
    if not text_value:
        return None
    parsed = pd.to_datetime(text_value, dayfirst=True, errors="coerce")
    if pd.isna(parsed):
        return None
    return pd.Timestamp(parsed).date()


def _explicit_date_from_time_text(value):
    """Extract an optional date written inside a time cell."""

    text_value = _clean_text(value)
    if not text_value:
        return None
    match = re.search(
        r"(?:\(|\s)(\d{1,2}[-/]\d{1,2}[-/]\d{2,4})(?:\)|\s|$)",
        text_value,
    )
    return _parse_log_date(match.group(1)) if match else None


def _clock_seconds(value):
    """Return seconds after midnight from an Excel/time-like value."""

    if value is None or pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        value = value.to_pydatetime()
    if isinstance(value, datetime):
        return (
            value.hour * 3600
            + value.minute * 60
            + value.second
            + value.microsecond / 1_000_000
        )
    if isinstance(value, time):
        return (
            value.hour * 3600
            + value.minute * 60
            + value.second
            + value.microsecond / 1_000_000
        )
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        numeric = float(value)
        return numeric * 24 * 3600 if 0 <= numeric < 1 else None

    text_value = _clean_text(value)
    if not text_value:
        return None
    match = re.search(
        r"(?<!\d)([01]?\d|2[0-3])[:.]([0-5]\d)(?::([0-5]\d))?",
        text_value,
    )
    if not match:
        return None
    return (
        int(match.group(1)) * 3600
        + int(match.group(2)) * 60
        + int(match.group(3) or 0)
    )


def _optional_distance_km(row):
    """Read a log distance when the workbook later supplies one."""

    for column in OPTIONAL_DISTANCE_COLUMNS:
        if column not in row.index:
            continue
        value = row.get(column)
        if value is None or pd.isna(value):
            continue
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return float(value)
        match = re.search(r"[-+]?\d+(?:\.\d+)?", str(value))
        if match:
            return float(match.group(0))
    return None


def _timestamp_text(value):
    """Return a CSV-stable timestamp string."""

    if value is None or pd.isna(value):
        return None
    return pd.Timestamp(value).isoformat(sep=" ")


def _log_record_id(record):
    """Create a deterministic, row-specific log identifier."""

    identity = "|".join(
        [
            str(record.get("log_excel_row") or ""),
            str(record.get("log_date") or ""),
            str(record.get("log_vehicle_type") or ""),
            str(record.get("log_start_location_key") or ""),
            str(record.get("log_end_location_key") or ""),
            str(record.get("log_start_time_original") or ""),
            str(record.get("log_end_time_original") or ""),
        ]
    )
    digest = hashlib.sha256(identity.encode("utf-8")).hexdigest()[:14]
    return f"log_{digest}"


def load_trip_log(trip_log_path, sheet_name=TRIP_LOG_SHEET_NAME):
    """Load and normalize the authoritative detailed trip-log sheet."""

    trip_log_path = Path(trip_log_path)
    if not trip_log_path.exists():
        raise FileNotFoundError(
            "Trip-log workbook was not found: "
            f"{trip_log_path}. Keep the workbook at this path or set "
            "reference.trip_log_path in config/settings.yaml."
        )
    try:
        dataframe = pd.read_excel(
            trip_log_path, sheet_name=sheet_name, engine="openpyxl"
        )
    except ValueError as error:
        raise ValueError(
            f"Worksheet '{sheet_name}' was not found in {trip_log_path.name}."
        ) from error

    dataframe.columns = [str(column).strip() for column in dataframe.columns]
    missing = sorted(REQUIRED_LOG_COLUMNS.difference(dataframe.columns))
    if missing:
        raise ValueError(
            "Trip-log detail sheet is missing required columns: "
            + ", ".join(missing)
        )

    records = []
    for dataframe_index, row in dataframe.iterrows():
        excel_row = int(dataframe_index) + 2
        log_date = _parse_log_date(row.get("Date"))
        start_location = _clean_text(row.get("Start Location"))
        end_location = _clean_text(row.get("End Location"))
        start_original = _clean_text(row.get("Start time"))
        end_original = _clean_text(row.get("End Time"))
        record = {
            "log_sheet": sheet_name,
            "log_excel_row": excel_row,
            "log_date": log_date,
            "log_day_number": _clean_text(row.get("Day No")),
            "log_vehicle_number": _clean_text(row.get("Vehicle Number")),
            "log_vehicle_type": _vehicle_type(row.get("Vehicle Type")),
            "log_driver_name": _clean_text(row.get("Driver Name")),
            "log_route_name": _clean_text(row.get("Route Name")),
            "log_start_location": start_location,
            "log_end_location": end_location,
            "log_start_location_key": canonical_location(start_location),
            "log_end_location_key": canonical_location(end_location),
            "log_start_time_original": start_original,
            "log_end_time_original": end_original,
            "log_start_clock_seconds": _clock_seconds(row.get("Start time")),
            "log_end_clock_seconds": _clock_seconds(row.get("End Time")),
            "log_start_explicit_date": _explicit_date_from_time_text(
                row.get("Start time")
            ),
            "log_end_explicit_date": _explicit_date_from_time_text(
                row.get("End Time")
            ),
            "log_file_name": _clean_text(row.get("File Name")),
            "log_collected_by": _clean_text(row.get("Collected by")),
            "log_charging_time": _clean_text(row.get("Charging time")),
            "log_battery_range": _clean_text(row.get("Battery Range")),
            "log_charging_type": _clean_text(row.get("Charging type")),
            "log_additional": _clean_text(row.get("Additional")),
            "log_distance_km": _optional_distance_km(row),
        }
        record["log_record_id"] = _log_record_id(record)

        if (
            log_date is None
            or record["log_vehicle_type"] == "UNKNOWN"
            or not record["log_start_location_key"]
            or not record["log_end_location_key"]
        ):
            record["log_record_status"] = "INVALID_LOG_IDENTITY"
        elif record["log_start_clock_seconds"] is None:
            record["log_record_status"] = "MISSING_LOG_START_TIME"
        elif record["log_end_clock_seconds"] is None:
            record["log_record_status"] = "MISSING_LOG_END_TIME"
        else:
            record["log_record_status"] = "VALID"
        records.append(record)

    result = pd.DataFrame(records)
    if result["log_record_id"].duplicated().any():
        raise ValueError("Trip log produced duplicate internal record IDs.")
    return result


def load_file_trip_index(index_path):
    """Load the completed file-level telemetry index."""

    index_path = Path(index_path)
    if not index_path.exists():
        raise FileNotFoundError(
            f"Trip index was not found: {index_path}. "
            "Run update_pipeline.bat first."
        )

    dataframe = None
    last_error = None
    for encoding in ["utf-8-sig", "utf-8", "cp1252"]:
        try:
            dataframe = pd.read_csv(index_path, encoding=encoding)
            break
        except UnicodeDecodeError as error:
            last_error = error
        except pd.errors.EmptyDataError as error:
            raise ValueError(f"Trip index is empty: {index_path}") from error
    if dataframe is None:
        raise ValueError(f"Trip index encoding could not be read: {last_error}")

    missing = sorted(REQUIRED_INDEX_COLUMNS.difference(dataframe.columns))
    if missing:
        raise ValueError(
            "Trip index is missing reconciliation fields: " + ", ".join(missing)
        )

    dataframe = dataframe.copy()
    dataframe["vehicle_type"] = dataframe["vehicle_type"].map(_vehicle_type)
    dataframe["start_time"] = pd.to_datetime(
        dataframe["start_time"], errors="coerce"
    )
    dataframe["end_time"] = pd.to_datetime(
        dataframe["end_time"], errors="coerce"
    )
    invalid = (
        dataframe["trip_id"].isna()
        | dataframe["cleaned_file_path"].isna()
        | dataframe["start_time"].isna()
        | dataframe["end_time"].isna()
        | (dataframe["vehicle_type"] == "UNKNOWN")
    )
    if bool(invalid.any()):
        files = ", ".join(
            dataframe.loc[invalid, "file_name"].astype(str).head(10).tolist()
        )
        raise ValueError(
            f"{int(invalid.sum())} indexed file(s) have invalid type, path, "
            f"or timestamps: {files}. Rerun update_pipeline.bat after "
            "correcting those inputs."
        )
    if dataframe["trip_id"].duplicated().any():
        raise ValueError("Trip index contains duplicate trip_id values.")
    return dataframe


def _location_similarity(first, second):
    """Return a conservative 0-1 similarity between canonical locations."""

    first_key = canonical_location(first)
    second_key = canonical_location(second)
    if not first_key or not second_key:
        return 0.0
    if first_key == second_key:
        return 1.0
    if min(len(first_key), len(second_key)) >= 5 and (
        first_key in second_key or second_key in first_key
    ):
        return 0.92
    return SequenceMatcher(None, first_key, second_key).ratio()


def _known_locations(log_records):
    locations = set()
    for column in ["log_start_location_key", "log_end_location_key"]:
        locations.update(
            value
            for value in log_records[column].dropna().astype(str)
            if value
        )
    return sorted(locations, key=lambda value: (-len(value), value))


def _route_occurrences(route_text, known_locations):
    compact_route = _ascii_compact(route_text)
    occurrences = []
    for location in known_locations:
        aliases = {
            alias
            for alias, canonical in LOCATION_ALIASES.items()
            if canonical == location
        }
        aliases.add(_ascii_compact(location))
        for search_value in sorted(
            (value for value in aliases if value),
            key=lambda value: (-len(value), value),
        ):
            start_position = 0
            while True:
                position = compact_route.find(search_value, start_position)
                if position < 0:
                    break
                occurrences.append((position, location))
                start_position = position + max(1, len(search_value))

    occurrences.sort(key=lambda item: (item[0], -len(item[1]), item[1]))
    deduplicated = []
    occupied = set()
    for position, location in occurrences:
        if position in occupied:
            continue
        deduplicated.append((position, location))
        occupied.add(position)
    return deduplicated


def _route_direction_score(indexed_trip, log_record, known_locations):
    """Score whether the filename carries the logged direction in order."""

    route_text = _clean_text(indexed_trip.get("file_name"))
    if not route_text:
        route_text = _clean_text(indexed_trip.get("route_name"))
    sequence = [
        location
        for _, location in _route_occurrences(route_text, known_locations)
    ]
    start_key = log_record.get("log_start_location_key", "")
    end_key = log_record.get("log_end_location_key", "")
    start_positions = [
        index
        for index, location in enumerate(sequence)
        if _location_similarity(location, start_key) >= 0.9
    ]
    end_positions = [
        index
        for index, location in enumerate(sequence)
        if _location_similarity(location, end_key) >= 0.9
    ]
    if start_positions and end_positions:
        forward = [
            (start, end)
            for start in start_positions
            for end in end_positions
            if start < end
        ]
        if forward:
            score = 1.0 if any(end == start + 1 for start, end in forward) else 0.9
        else:
            score = 0.2
    elif start_positions or end_positions:
        score = 0.45
    else:
        score = 0.0
    return score, " > ".join(sequence)


def _date_difference_days(log_date, indexed_trip):
    """Use the input date folder as the primary collection-date evidence."""

    index_date = _parse_log_date(indexed_trip.get("date_folder"))
    if index_date is None:
        relative = str(indexed_trip.get("relative_path", "")).replace("/", "\\")
        index_date = _parse_log_date(relative.split("\\")[0])
    if index_date is None:
        index_date = pd.Timestamp(indexed_trip["start_time"]).date()
    return abs((log_date - index_date).days)


def _distance_to_interval_seconds(timestamp, interval_start, interval_end):
    timestamp = pd.Timestamp(timestamp)
    interval_start = pd.Timestamp(interval_start)
    interval_end = pd.Timestamp(interval_end)
    if timestamp < interval_start:
        return (interval_start - timestamp).total_seconds()
    if timestamp > interval_end:
        return (timestamp - interval_end).total_seconds()
    return 0.0


def _datetime_candidates(log_date, clock_seconds, explicit_date=None):
    """Generate AM/PM and nearby-day interpretations of a manual clock."""

    if log_date is None or clock_seconds is None or pd.isna(clock_seconds):
        return []
    base_date = explicit_date or log_date
    # Manual field logs can occasionally contain a 12-hour AM/PM
    # transcription offset. Test both the written clock and its 12-hour
    # counterpart, then let the actual Packet-Time telemetry select the
    # interpretation with real sample support.
    clock_value = float(clock_seconds) % 86_400.0
    clock_variants = {
        clock_value,
        (clock_value + 12 * 3600.0) % 86_400.0,
    }

    day_offsets = [0] if explicit_date else [-1, 0, 1]
    candidates = set()
    for day_offset in day_offsets:
        day_value = base_date + timedelta(days=day_offset)
        midnight = datetime.combine(day_value, time.min)
        for variant_seconds in clock_variants:
            candidates.add(
                pd.Timestamp(
                    midnight + timedelta(seconds=variant_seconds)
                )
            )
    return sorted(candidates)


def _nearest_sample_seconds(timestamps, target):
    """Return absolute seconds to the nearest actual GPS timestamp."""

    if target is None or len(timestamps) == 0:
        return float("inf")
    target = pd.Timestamp(target)
    position = int(timestamps.searchsorted(target, side="left"))
    offsets = []
    if position < len(timestamps):
        offsets.append(abs((pd.Timestamp(timestamps.iloc[position]) - target).total_seconds()))
    if position > 0:
        offsets.append(abs((pd.Timestamp(timestamps.iloc[position - 1]) - target).total_seconds()))
    return min(offsets) if offsets else float("inf")


def _resolve_log_window(
    log_record,
    raw_start,
    raw_end,
    timestamps=None,
    maximum_duration_hours=30.0,
):
    """Choose the AM/PM/overnight interpretation supported by telemetry."""

    raw_start = pd.Timestamp(raw_start)
    raw_end = pd.Timestamp(raw_end)
    start_candidates = _datetime_candidates(
        log_record.get("log_date"),
        log_record.get("log_start_clock_seconds"),
        log_record.get("log_start_explicit_date"),
    )
    end_candidates = _datetime_candidates(
        log_record.get("log_date"),
        log_record.get("log_end_clock_seconds"),
        log_record.get("log_end_explicit_date"),
    )
    if not start_candidates:
        return None, None

    if timestamps is not None:
        timestamps = pd.Series(pd.to_datetime(timestamps)).sort_values().reset_index(drop=True)

    def start_key(candidate):
        if timestamps is not None:
            sample_offset = _nearest_sample_seconds(timestamps, candidate)
        else:
            sample_offset = _distance_to_interval_seconds(candidate, raw_start, raw_end)
        expected_date = log_record.get("log_start_explicit_date") or log_record.get("log_date")
        date_penalty = abs((candidate.date() - expected_date).days) * 86_400
        return (sample_offset + date_penalty, abs((candidate - raw_start).total_seconds()))

    if not end_candidates:
        return min(start_candidates, key=start_key), None

    windows = []
    for start_candidate in start_candidates:
        for end_candidate in end_candidates:
            duration_hours = (end_candidate - start_candidate).total_seconds() / 3600.0
            if duration_hours <= 0 or duration_hours > float(maximum_duration_hours):
                continue
            if timestamps is not None:
                left = int(timestamps.searchsorted(start_candidate, side="left"))
                right = int(timestamps.searchsorted(end_candidate, side="right"))
                points = max(0, right - left)
                start_offset = _nearest_sample_seconds(timestamps, start_candidate)
                end_offset = _nearest_sample_seconds(timestamps, end_candidate)
                no_points_penalty = 7 * 86_400 if points == 0 else 0
                sample_cost = start_offset + end_offset + no_points_penalty
            else:
                points = 0
                sample_cost = (
                    _distance_to_interval_seconds(start_candidate, raw_start, raw_end)
                    + _distance_to_interval_seconds(end_candidate, raw_start, raw_end)
                )
            expected_date = log_record.get("log_start_explicit_date") or log_record.get("log_date")
            date_penalty = abs((start_candidate.date() - expected_date).days) * 86_400
            windows.append(
                (
                    sample_cost + date_penalty,
                    -points,
                    abs((start_candidate - raw_start).total_seconds()),
                    start_candidate,
                    end_candidate,
                )
            )
    if not windows:
        return min(start_candidates, key=start_key), None
    best = min(windows, key=lambda item: item[:3])
    return best[3], best[4]


def _circular_clock_difference_minutes(clock_seconds, timestamp):
    if clock_seconds is None or timestamp is None or pd.isna(timestamp):
        return None
    timestamp = pd.Timestamp(timestamp)
    sample_seconds = (
        timestamp.hour * 3600
        + timestamp.minute * 60
        + timestamp.second
        + timestamp.microsecond / 1_000_000
    )
    direct = abs(float(clock_seconds) - sample_seconds)
    return min(direct, 86_400 - direct) / 60.0


def _evaluate_match(indexed_trip, log_record, known_locations, rules):
    """Score one possible file/log assignment without using a hard time gate."""

    if indexed_trip["vehicle_type"] != log_record["log_vehicle_type"]:
        return None
    log_date = log_record.get("log_date")
    if log_date is None:
        return None
    date_difference = _date_difference_days(log_date, indexed_trip)
    if date_difference > int(rules["maximum_log_date_difference_days"]):
        return None

    route_score, route_sequence = _route_direction_score(
        indexed_trip, log_record, known_locations
    )
    log_file_key = _ascii_compact(log_record.get("log_file_name"))
    indexed_file_key = _ascii_compact(indexed_trip.get("file_name"))
    exact_logged_filename = bool(log_file_key and log_file_key == indexed_file_key)
    if (
        route_score < float(rules["minimum_route_direction_score"])
        and not exact_logged_filename
    ):
        return None

    raw_start = pd.Timestamp(indexed_trip["start_time"])
    raw_end = pd.Timestamp(indexed_trip["end_time"])
    resolved_start, resolved_end = _resolve_log_window(
        log_record,
        raw_start,
        raw_end,
        maximum_duration_hours=rules["maximum_verified_duration_hours"],
    )
    start_clock_difference = _circular_clock_difference_minutes(
        log_record.get("log_start_clock_seconds"), raw_start
    )
    end_clock_difference = _circular_clock_difference_minutes(
        log_record.get("log_end_clock_seconds"), raw_end
    )
    available_clock_differences = [
        value
        for value in [start_clock_difference, end_clock_difference]
        if value is not None
    ]
    mean_clock_difference = (
        sum(available_clock_differences) / len(available_clock_differences)
        if available_clock_differences
        else 720.0
    )

    detected = [value for value in route_sequence.split(" > ") if value]
    exact_direction = detected == [
        log_record.get("log_start_location_key"),
        log_record.get("log_end_location_key"),
    ]
    filename_quality = 5.0 if exact_direction else 2.0
    if any(token in indexed_file_key for token in ["packetreport", "download", "zpa"]):
        filename_quality = max(0.0, filename_quality - 2.0)

    # Date, vehicle type, and ordered route dominate.  Clock evidence resolves
    # repeated same-route trips but cannot exclude files containing extra data.
    date_points = 60.0 if date_difference == 0 else 15.0
    route_points = 60.0 * route_score
    time_points = max(0.0, 20.0 - mean_clock_difference / 36.0)
    logged_filename_points = 100.0 if exact_logged_filename else 0.0
    score = (
        date_points
        + route_points
        + time_points
        + filename_quality
        + logged_filename_points
    )
    return {
        "source_trip_id": indexed_trip["trip_id"],
        "log_record_id": log_record["log_record_id"],
        "match_score": round(score, 3),
        "date_difference_days": int(date_difference),
        "route_direction_score": round(route_score, 3),
        "detected_file_route_sequence": route_sequence,
        "resolved_log_start": resolved_start,
        "resolved_log_end": resolved_end,
        "start_clock_difference_minutes": (
            round(start_clock_difference, 3)
            if start_clock_difference is not None
            else None
        ),
        "end_clock_difference_minutes": (
            round(end_clock_difference, 3)
            if end_clock_difference is not None
            else None
        ),
        "exact_logged_filename_match": exact_logged_filename,
    }


def _hungarian_minimize(cost_matrix):
    """Pure-Python rectangular Hungarian assignment (rows <= columns)."""

    row_count = len(cost_matrix)
    column_count = len(cost_matrix[0]) if row_count else 0
    if row_count == 0:
        return []
    if row_count > column_count:
        raise ValueError("Hungarian assignment requires rows <= columns.")

    u = [0.0] * (row_count + 1)
    v = [0.0] * (column_count + 1)
    p = [0] * (column_count + 1)
    way = [0] * (column_count + 1)
    for i in range(1, row_count + 1):
        p[0] = i
        j0 = 0
        minimum = [float("inf")] * (column_count + 1)
        used = [False] * (column_count + 1)
        while True:
            used[j0] = True
            i0 = p[j0]
            delta = float("inf")
            j1 = 0
            for j in range(1, column_count + 1):
                if used[j]:
                    continue
                current = cost_matrix[i0 - 1][j - 1] - u[i0] - v[j]
                if current < minimum[j]:
                    minimum[j] = current
                    way[j] = j0
                if minimum[j] < delta:
                    delta = minimum[j]
                    j1 = j
            for j in range(column_count + 1):
                if used[j]:
                    u[p[j]] += delta
                    v[j] -= delta
                else:
                    minimum[j] -= delta
            j0 = j1
            if p[j0] == 0:
                break
        while True:
            j1 = way[j0]
            p[j0] = p[j1]
            j0 = j1
            if j0 == 0:
                break

    assignment = [-1] * row_count
    for j in range(1, column_count + 1):
        if p[j] != 0:
            assignment[p[j] - 1] = j - 1
    return assignment


def match_log_rows_to_indexed_files(indexed_trips, log_records, rules):
    """Globally assign each file to at most one log row and vice versa."""

    logs = log_records.sort_values("log_excel_row", kind="stable").to_dict(
        orient="records"
    )
    files = indexed_trips.sort_values(
        ["relative_path", "trip_id"], kind="stable"
    ).to_dict(orient="records")
    known_locations = _known_locations(log_records)
    candidates = {}
    candidates_by_log = {record["log_record_id"]: [] for record in logs}
    for log_record in logs:
        for indexed_trip in files:
            candidate = _evaluate_match(
                indexed_trip, log_record, known_locations, rules
            )
            if candidate is None:
                continue
            key = (log_record["log_record_id"], indexed_trip["trip_id"])
            candidates[key] = candidate
            candidates_by_log[log_record["log_record_id"]].append(candidate)

    minimum_score = float(rules["minimum_match_score"])
    invalid_score = -1_000_000.0
    score_matrix = []
    for log_record in logs:
        row = []
        for indexed_trip in files:
            candidate = candidates.get(
                (log_record["log_record_id"], indexed_trip["trip_id"])
            )
            row.append(candidate["match_score"] if candidate else invalid_score)
        row.extend([0.0] * len(logs))  # independent unmatched dummies
        score_matrix.append(row)

    maximum_score = max(
        [minimum_score]
        + [candidate["match_score"] for candidate in candidates.values()]
    )
    cost_matrix = [
        [maximum_score - score for score in row] for row in score_matrix
    ]
    assignment = _hungarian_minimize(cost_matrix)

    assigned_file_to_log = {}
    for row_index, column_index in enumerate(assignment):
        if not (0 <= column_index < len(files)):
            continue
        log_record = logs[row_index]
        indexed_trip = files[column_index]
        candidate = candidates.get(
            (log_record["log_record_id"], indexed_trip["trip_id"])
        )
        if candidate and candidate["match_score"] >= minimum_score:
            assigned_file_to_log[indexed_trip["trip_id"]] = log_record[
                "log_record_id"
            ]

    selected = []
    unmatched = []
    for row_index, column_index in enumerate(assignment):
        log_record = logs[row_index]
        candidate = None
        if 0 <= column_index < len(files):
            indexed_trip = files[column_index]
            candidate = candidates.get(
                (log_record["log_record_id"], indexed_trip["trip_id"])
            )
        if candidate and candidate["match_score"] >= minimum_score:
            alternatives = sorted(
                candidates_by_log[log_record["log_record_id"]],
                key=lambda item: (-item["match_score"], item["source_trip_id"]),
            )
            second_scores = [
                item["match_score"]
                for item in alternatives
                if item["source_trip_id"] != candidate["source_trip_id"]
            ]
            margin = (
                candidate["match_score"] - max(second_scores)
                if second_scores
                else candidate["match_score"]
            )
            candidate = dict(candidate)
            candidate["one_to_one_assignment"] = True
            candidate["assignment_margin"] = round(margin, 3)
            candidate["assignment_confidence"] = (
                "HIGH"
                if margin >= float(rules["high_confidence_assignment_margin"])
                else "MEDIUM"
                if margin >= float(rules["low_confidence_assignment_margin"])
                else "LOW"
            )
            selected.append(candidate)
            continue

        record = dict(log_record)
        possible = sorted(
            candidates_by_log[log_record["log_record_id"]],
            key=lambda item: (-item["match_score"], item["source_trip_id"]),
        )
        if not possible:
            record["unmatched_reason"] = "NO_SAME_DATE_TYPE_DIRECTION_FILE"
            record["best_match_score"] = None
        else:
            best = possible[0]
            record["best_match_score"] = best["match_score"]
            record["best_source_trip_id"] = best["source_trip_id"]
            assigned_elsewhere = best["source_trip_id"] in assigned_file_to_log
            record["unmatched_reason"] = (
                "ONE_TO_ONE_CONFLICT_FILE_ASSIGNED_TO_ANOTHER_LOG_ROW"
                if assigned_elsewhere
                else "BEST_CANDIDATE_BELOW_MINIMUM_SCORE"
            )
        unmatched.append(record)

    selected_frame = pd.DataFrame(selected)
    if not selected_frame.empty:
        selected_frame = selected_frame.sort_values(
            "log_record_id", kind="stable"
        ).reset_index(drop=True)
    return selected_frame, pd.DataFrame(unmatched)


def _verified_trip_id(source_trip_id, log_record):
    identity = "|".join(
        [
            str(source_trip_id),
            str(log_record.get("log_record_id") or ""),
            str(log_record.get("log_date") or ""),
            str(log_record.get("log_start_location_key") or ""),
            str(log_record.get("log_end_location_key") or ""),
        ]
    )
    digest = hashlib.sha256(identity.encode("utf-8")).hexdigest()[:16]
    return f"verified_{digest}"


def _empty_reconciliation_record(indexed_trip, log_record, match_record):
    return {
        "verified_trip_id": _verified_trip_id(indexed_trip["trip_id"], log_record),
        "reconciliation_status": None,
        "reconciliation_reason": None,
        "boundary_method": None,
        "boundary_confidence": None,
        "boundary_clock_shift_minutes": None,
        "source_trip_id": indexed_trip["trip_id"],
        "source_file_name": indexed_trip["file_name"],
        "source_relative_path": indexed_trip["relative_path"],
        "source_cleaned_file_path": indexed_trip["cleaned_file_path"],
        "vehicle_type": indexed_trip["vehicle_type"],
        "file_route_name_provisional": indexed_trip.get("route_name"),
        "raw_first_timestamp": _timestamp_text(indexed_trip["start_time"]),
        "raw_last_timestamp": _timestamp_text(indexed_trip["end_time"]),
        "raw_first_latitude": indexed_trip.get("start_latitude"),
        "raw_first_longitude": indexed_trip.get("start_longitude"),
        "raw_last_latitude": indexed_trip.get("end_latitude"),
        "raw_last_longitude": indexed_trip.get("end_longitude"),
        "raw_index_distance_km": indexed_trip.get("distance_km"),
        "raw_index_points": indexed_trip.get("telemetry_points"),
        "log_record_id": log_record["log_record_id"],
        "log_excel_row": log_record["log_excel_row"],
        "log_date": str(log_record.get("log_date") or ""),
        "log_vehicle_number": log_record.get("log_vehicle_number"),
        "log_route_name": log_record.get("log_route_name"),
        "log_start_location": log_record.get("log_start_location"),
        "log_end_location": log_record.get("log_end_location"),
        "log_start_time_original": log_record.get("log_start_time_original"),
        "log_end_time_original": log_record.get("log_end_time_original"),
        "log_distance_km": log_record.get("log_distance_km"),
        "resolved_log_start": _timestamp_text(match_record.get("resolved_log_start")),
        "resolved_log_end": _timestamp_text(match_record.get("resolved_log_end")),
        "log_collected_by": log_record.get("log_collected_by"),
        "log_charging_time": log_record.get("log_charging_time"),
        "log_battery_range": log_record.get("log_battery_range"),
        "log_charging_type": log_record.get("log_charging_type"),
        "log_additional": log_record.get("log_additional"),
        "match_score": match_record.get("match_score"),
        "match_date_difference_days": match_record.get("date_difference_days"),
        "match_route_direction_score": match_record.get("route_direction_score"),
        "detected_file_route_sequence": match_record.get(
            "detected_file_route_sequence"
        ),
        "one_to_one_assignment": True,
        "assignment_confidence": match_record.get("assignment_confidence"),
        "assignment_margin": match_record.get("assignment_margin"),
        "verified_start_timestamp": None,
        "verified_end_timestamp": None,
        "verified_start_latitude": None,
        "verified_start_longitude": None,
        "verified_end_latitude": None,
        "verified_end_longitude": None,
        "verified_duration_minutes": None,
        "verified_distance_km": None,
        "verified_straight_line_distance_km": None,
        "verified_points": None,
        "start_sample_offset_seconds": None,
        "end_sample_offset_seconds": None,
        "points_removed_before": None,
        "points_removed_after": None,
        "minutes_removed_before": None,
        "minutes_removed_after": None,
        "distance_removed_before_km": None,
        "distance_removed_after_km": None,
        "raw_recalculated_distance_km": None,
        "distance_accounting_difference_km": None,
        "verified_minus_log_distance_km": None,
        "source_file_assignment_count": 1,
    }


def _consecutive_distances_km(telemetry):
    distances = [0.0] * len(telemetry)
    for index in range(1, len(telemetry)):
        distances[index] = haversine_distance_km(
            telemetry.iloc[index - 1]["Latitude"],
            telemetry.iloc[index - 1]["Longitude"],
            telemetry.iloc[index]["Latitude"],
            telemetry.iloc[index]["Longitude"],
        )
    return distances


def _activity_blocks(telemetry, rules):
    """Find continuous movement blocks while preserving stops inside a trip."""

    if len(telemetry) < 2:
        return []
    timestamps = telemetry["Timestamp"].reset_index(drop=True)
    distances = _consecutive_distances_km(telemetry)
    speed = pd.to_numeric(telemetry["Speed(km/h)"], errors="coerce").fillna(0.0)
    continuity_gap = float(rules["continuity_gap_minutes"]) * 60.0
    hard_gap = float(rules["hard_gap_minutes"]) * 60.0
    maximum_speed = float(rules["maximum_implied_speed_kmh"])
    breaks = [0]
    for index in range(1, len(telemetry)):
        seconds = (timestamps.iloc[index] - timestamps.iloc[index - 1]).total_seconds()
        distance = distances[index]

        # Timestamp is the cleaned telemetry timestamp derived from the source
        # Packet Time. Actual Received Time is not used for trip-boundary logic.
        #
        # A repeated Packet Time (seconds == 0) can occur when the device emits
        # duplicate/near-duplicate records. Do not split one continuous trip
        # merely because two consecutive packets share the same Packet Time.
        # A true backward time jump (seconds < 0) is still treated as a break.
        if seconds < 0:
            should_break = True
            implied = float("inf")
        elif seconds == 0:
            should_break = False
            implied = 0.0
        else:
            implied = distance / seconds * 3600.0
            should_break = False
            should_break |= seconds > hard_gap
            should_break |= seconds > continuity_gap and distance > 0.2
            should_break |= implied > maximum_speed and distance > 0.5

        if should_break:
            breaks.append(index)
    breaks.append(len(telemetry))

    blocks = []
    movement_speed = float(rules["movement_speed_threshold_kmh"])
    movement_step = float(rules["movement_step_metres"]) / 1000.0
    for group_number in range(len(breaks) - 1):
        group_start = breaks[group_number]
        group_end = breaks[group_number + 1] - 1
        if group_end <= group_start:
            continue
        active = []
        for index in range(group_start, group_end + 1):
            adjacent_distance = distances[index]
            if index + 1 < len(distances):
                adjacent_distance = max(adjacent_distance, distances[index + 1])
            if speed.iloc[index] >= movement_speed or adjacent_distance >= movement_step:
                active.append(index)
        if not active:
            continue
        start_index = active[0]
        end_index = active[-1]
        extension_seconds = (
            float(rules["boundary_stationary_extension_minutes"]) * 60.0
        )
        if active[0] > group_start:
            previous_gap = (
                timestamps.iloc[active[0]] - timestamps.iloc[active[0] - 1]
            ).total_seconds()
            if 0 < previous_gap <= extension_seconds:
                start_index = active[0] - 1
        if active[-1] < group_end:
            following_gap = (
                timestamps.iloc[active[-1] + 1] - timestamps.iloc[active[-1]]
            ).total_seconds()
            if 0 < following_gap <= extension_seconds:
                end_index = active[-1] + 1
        block = telemetry.iloc[start_index : end_index + 1]
        duration = (
            pd.Timestamp(block.iloc[-1]["Timestamp"])
            - pd.Timestamp(block.iloc[0]["Timestamp"])
        ).total_seconds() / 60.0
        if duration <= 0:
            continue
        distance = calculate_trip_distance_km(block)
        if distance < float(rules["minimum_verified_distance_km"]):
            continue
        blocks.append(
            {
                "start_index": start_index,
                "end_index": end_index,
                "start": pd.Timestamp(block.iloc[0]["Timestamp"]),
                "end": pd.Timestamp(block.iloc[-1]["Timestamp"]),
                "duration_minutes": duration,
                "distance_km": distance,
                "points": len(block),
            }
        )
    return blocks


def _choose_activity_block(blocks, resolved_start, resolved_end, log_distance, rules):
    """Choose the movement block most consistent with the intended trip."""

    if not blocks:
        return None, None, None
    expected_duration = None
    if resolved_start is not None and resolved_end is not None:
        expected_duration = (resolved_end - resolved_start).total_seconds() / 60.0

    scored = []
    for block in blocks:
        start_shift = (
            (block["start"] - resolved_start).total_seconds() / 60.0
            if resolved_start is not None
            else 0.0
        )
        end_shift = (
            (block["end"] - resolved_end).total_seconds() / 60.0
            if resolved_end is not None
            else start_shift
        )
        shift_consistency = abs(start_shift - end_shift)
        mean_shift = (start_shift + end_shift) / 2.0
        duration_error_ratio = (
            abs(block["duration_minutes"] - expected_duration)
            / max(expected_duration, 30.0)
            if expected_duration is not None
            else 0.0
        )
        distance_error_ratio = (
            abs(block["distance_km"] - float(log_distance))
            / max(float(log_distance), 1.0)
            if log_distance is not None and not pd.isna(log_distance)
            else 0.0
        )
        score = 100.0
        score -= min(duration_error_ratio, 3.0) * 30.0
        score -= min(shift_consistency, 240.0) * 0.15
        score -= min(abs(mean_shift), 720.0) * 0.02
        score -= min(distance_error_ratio, 3.0) * 25.0
        score += min(block["distance_km"], 200.0) * 0.05
        scored.append((score, block, mean_shift, shift_consistency))
    scored.sort(
        key=lambda item: (-item[0], -item[1]["distance_km"], item[1]["start"])
    )
    best = scored[0]
    margin = best[0] - scored[1][0] if len(scored) > 1 else 100.0
    return best, margin, scored


def _verify_one_trip(indexed_trip, log_record, match_record, telemetry, rules):
    """Select one safe boundary pair from the full telemetry series."""

    result = _empty_reconciliation_record(indexed_trip, log_record, match_record)
    if log_record["log_record_status"] == "INVALID_LOG_IDENTITY":
        result["reconciliation_status"] = "DATA_CHECK_REQUIRED"
        result["reconciliation_reason"] = (
            "Trip-log date, vehicle type, or route endpoint is missing"
        )
        return result

    timestamps = telemetry["Timestamp"].reset_index(drop=True)
    resolved_start, resolved_end = _resolve_log_window(
        log_record,
        timestamps.iloc[0],
        timestamps.iloc[-1],
        timestamps=timestamps,
        maximum_duration_hours=rules["maximum_verified_duration_hours"],
    )
    result["resolved_log_start"] = _timestamp_text(resolved_start)
    result["resolved_log_end"] = _timestamp_text(resolved_end)
    if resolved_start is None:
        result["reconciliation_status"] = "DATA_CHECK_REQUIRED"
        result["reconciliation_reason"] = "No usable trip-log start time"
        return result

    start_position = int(timestamps.searchsorted(resolved_start, side="left"))
    end_position = (
        int(timestamps.searchsorted(resolved_end, side="right")) - 1
        if resolved_end is not None
        else -1
    )
    exact_available = (
        resolved_end is not None
        and 0 <= start_position < len(telemetry)
        and 0 <= end_position < len(telemetry)
        and start_position <= end_position
    )
    maximum_offset = float(rules["maximum_boundary_sample_offset_minutes"]) * 60.0
    if exact_available:
        start_offset = abs(
            (pd.Timestamp(timestamps.iloc[start_position]) - resolved_start).total_seconds()
        )
        end_offset = abs(
            (resolved_end - pd.Timestamp(timestamps.iloc[end_position])).total_seconds()
        )
        exact = telemetry.iloc[start_position : end_position + 1]
        exact_distance = calculate_trip_distance_km(exact)
        exact_available = (
            len(exact) >= int(rules["minimum_verified_points"])
            and start_offset <= maximum_offset
            and end_offset <= maximum_offset
            and exact_distance >= float(rules["minimum_verified_distance_km"])
        )
    else:
        start_offset = None
        end_offset = None

    if exact_available:
        boundary_start = start_position
        boundary_end = end_position
        method = "LOG_ANCHORED"
        confidence = "HIGH"
        clock_shift = 0.0
        reason = "Full telemetry supports the logged start/end window"
    else:
        blocks = _activity_blocks(telemetry, rules)
        best, margin, _ = _choose_activity_block(
            blocks,
            resolved_start,
            resolved_end,
            log_record.get("log_distance_km"),
            rules,
        )
        if best is None or best[0] < float(rules["minimum_activity_block_score"]):
            result["reconciliation_status"] = "DATA_CHECK_REQUIRED"
            result["reconciliation_reason"] = (
                "No unique movement block supports the intended trip boundaries"
            )
            result["start_sample_offset_seconds"] = start_offset
            result["end_sample_offset_seconds"] = end_offset
            return result
        score, block, mean_shift, shift_consistency = best
        confidence = (
            "HIGH"
            if score >= 75.0 and margin >= 10.0 and shift_consistency <= 30.0
            else "MEDIUM"
            if score >= 55.0 and margin >= float(rules["minimum_activity_score_margin"])
            else "LOW"
        )
        if confidence == "LOW":
            result["reconciliation_status"] = "DATA_CHECK_REQUIRED"
            result["reconciliation_reason"] = (
                "Several movement blocks fit the log similarly; boundary not forced"
            )
            result["boundary_method"] = "ACTIVITY_BLOCK_AMBIGUOUS"
            result["boundary_confidence"] = confidence
            return result
        boundary_start = block["start_index"]
        boundary_end = block["end_index"]
        method = (
            "LOG_START_ACTIVITY_END"
            if resolved_end is None
            else "ACTIVITY_INFERRED"
        )
        clock_shift = mean_shift
        reason = (
            "Movement block selected using log time/duration and telemetry distance"
            if resolved_end is not None
            else "Missing log end inferred from the movement block after logged start"
        )

    verified = telemetry.iloc[boundary_start : boundary_end + 1].copy()
    if len(verified) < int(rules["minimum_verified_points"]):
        result["reconciliation_status"] = "DATA_CHECK_REQUIRED"
        result["reconciliation_reason"] = "Verified boundary has too few GPS points"
        return result
    verified_start = pd.Timestamp(verified.iloc[0]["Timestamp"])
    verified_end = pd.Timestamp(verified.iloc[-1]["Timestamp"])
    duration_minutes = (verified_end - verified_start).total_seconds() / 60.0
    if duration_minutes <= 0 or duration_minutes > float(
        rules["maximum_verified_duration_hours"]
    ) * 60.0:
        result["reconciliation_status"] = "DATA_CHECK_REQUIRED"
        result["reconciliation_reason"] = "Verified trip duration is invalid"
        return result

    raw_distance = calculate_trip_distance_km(telemetry)
    verified_distance = calculate_trip_distance_km(verified)
    prefix = telemetry.iloc[: boundary_start + 1]
    suffix = telemetry.iloc[boundary_end:]
    removed_before = calculate_trip_distance_km(prefix)
    removed_after = calculate_trip_distance_km(suffix)
    accounting = raw_distance - (
        removed_before + verified_distance + removed_after
    )
    straight_line = haversine_distance_km(
        verified.iloc[0]["Latitude"],
        verified.iloc[0]["Longitude"],
        verified.iloc[-1]["Latitude"],
        verified.iloc[-1]["Longitude"],
    )
    log_distance = log_record.get("log_distance_km")
    distance_difference = (
        verified_distance - float(log_distance)
        if log_distance is not None and not pd.isna(log_distance)
        else None
    )
    result.update(
        {
            "reconciliation_status": "READY_FOR_PAIRING",
            "reconciliation_reason": reason,
            "boundary_method": method,
            "boundary_confidence": confidence,
            "boundary_clock_shift_minutes": round(clock_shift, 3),
            "verified_start_timestamp": _timestamp_text(verified_start),
            "verified_end_timestamp": _timestamp_text(verified_end),
            "verified_start_latitude": round(float(verified.iloc[0]["Latitude"]), 6),
            "verified_start_longitude": round(float(verified.iloc[0]["Longitude"]), 6),
            "verified_end_latitude": round(float(verified.iloc[-1]["Latitude"]), 6),
            "verified_end_longitude": round(float(verified.iloc[-1]["Longitude"]), 6),
            "verified_duration_minutes": round(duration_minutes, 3),
            "verified_distance_km": round(verified_distance, 6),
            "verified_straight_line_distance_km": round(straight_line, 6),
            "verified_points": int(len(verified)),
            "start_sample_offset_seconds": (
                round(abs((verified_start - resolved_start).total_seconds()), 3)
                if resolved_start is not None
                else None
            ),
            "end_sample_offset_seconds": (
                round(abs((resolved_end - verified_end).total_seconds()), 3)
                if resolved_end is not None
                else None
            ),
            "points_removed_before": int(boundary_start),
            "points_removed_after": int(len(telemetry) - boundary_end - 1),
            "minutes_removed_before": round(
                (verified_start - pd.Timestamp(telemetry.iloc[0]["Timestamp"])).total_seconds()
                / 60.0,
                3,
            ),
            "minutes_removed_after": round(
                (pd.Timestamp(telemetry.iloc[-1]["Timestamp"]) - verified_end).total_seconds()
                / 60.0,
                3,
            ),
            "distance_removed_before_km": round(removed_before, 6),
            "distance_removed_after_km": round(removed_after, 6),
            "raw_recalculated_distance_km": round(raw_distance, 6),
            "distance_accounting_difference_km": round(accounting, 9),
            "verified_minus_log_distance_km": (
                round(distance_difference, 6)
                if distance_difference is not None
                else None
            ),
        }
    )
    return result


def build_verified_trip_records(indexed_trips, log_records, selected_matches, rules):
    """Verify every unique match; never reuse or split a source file."""

    if selected_matches.empty:
        return pd.DataFrame()
    if selected_matches["source_trip_id"].duplicated().any():
        raise RuntimeError("One-to-one protection failed: a source file was reused.")
    if selected_matches["log_record_id"].duplicated().any():
        raise RuntimeError("One-to-one protection failed: a log row was reused.")

    indexed_lookup = {
        record["trip_id"]: record
        for record in indexed_trips.to_dict(orient="records")
    }
    log_lookup = {
        record["log_record_id"]: record
        for record in log_records.to_dict(orient="records")
    }
    records = []
    for match_record in selected_matches.to_dict(orient="records"):
        indexed_trip = indexed_lookup[match_record["source_trip_id"]]
        log_record = log_lookup[match_record["log_record_id"]]
        try:
            telemetry = load_cleaned_telemetry(indexed_trip["cleaned_file_path"])
            record = _verify_one_trip(
                indexed_trip, log_record, match_record, telemetry, rules
            )
        except Exception as error:
            record = _empty_reconciliation_record(
                indexed_trip, log_record, match_record
            )
            record["reconciliation_status"] = "DATA_CHECK_REQUIRED"
            record["reconciliation_reason"] = (
                f"Cleaned telemetry could not be verified: {type(error).__name__}: {error}"
            )
        records.append(record)

    result = pd.DataFrame(records)
    if not result.empty:
        result = result.sort_values(
            ["log_date", "log_excel_row", "source_file_name"],
            kind="stable",
        ).reset_index(drop=True)
    return result


def _trip_log_path_from_config(config):
    configured = config.get("reference", {}).get("trip_log_path")
    return Path(configured) if configured else DEFAULT_TRIP_LOG_PATH


def _normalise_relative_path(value):
    """Return a stable case-insensitive relative-path comparison key."""

    text_value = str(value or "").strip().replace("\\", "/")
    while text_value.startswith("./"):
        text_value = text_value[2:]
    return text_value.casefold()


def _discover_current_input_files(config):
    """Discover current raw telemetry workbooks without depending on Git state."""

    input_root = Path(config["project"]["input_folder"])
    if not input_root.exists():
        raise FileNotFoundError(f"Configured input folder was not found: {input_root}")
    if not input_root.is_dir():
        raise NotADirectoryError(f"Configured input folder is not a directory: {input_root}")

    configured_extensions = config.get("files", {}).get(
        "accepted_extensions", [".xlsx"]
    )
    extensions = {
        str(extension).strip().casefold()
        if str(extension).strip().startswith(".")
        else "." + str(extension).strip().casefold()
        for extension in configured_extensions
        if str(extension).strip()
    }
    if not extensions:
        extensions = {".xlsx"}

    records = []
    for path in sorted(
        input_root.rglob("*"),
        key=lambda item: str(item).casefold(),
    ):
        if not path.is_file():
            continue
        if path.name.startswith("~$"):
            continue
        if path.suffix.casefold() not in extensions:
            continue
        records.append(
            {
                "input_file_path": str(path.resolve()),
                "input_relative_path": str(path.relative_to(input_root)),
                "input_relative_key": _normalise_relative_path(
                    path.relative_to(input_root)
                ),
            }
        )
    return pd.DataFrame(records)


def _check_trip_index_freshness(config, indexed_trips):
    """Compare current Data/Input workbooks with the persisted trip index."""

    input_files = _discover_current_input_files(config)
    if input_files.empty:
        return {
            "input_file_count": 0,
            "indexed_file_count": len(indexed_trips),
            "raw_not_indexed": pd.DataFrame(),
            "indexed_without_raw": pd.DataFrame(),
        }

    indexed_keys = {
        _normalise_relative_path(value)
        for value in indexed_trips.get("relative_path", pd.Series(dtype=str))
        .dropna()
        .astype(str)
    }
    raw_keys = set(input_files["input_relative_key"].astype(str))

    raw_not_indexed = input_files.loc[
        ~input_files["input_relative_key"].isin(indexed_keys)
    ].copy()

    indexed_without_raw = indexed_trips.loc[
        ~indexed_trips["relative_path"]
        .astype(str)
        .map(_normalise_relative_path)
        .isin(raw_keys)
    ].copy()

    return {
        "input_file_count": len(input_files),
        "indexed_file_count": len(indexed_trips),
        "raw_not_indexed": raw_not_indexed,
        "indexed_without_raw": indexed_without_raw,
    }


def _format_stale_index_error(freshness):
    missing = freshness["raw_not_indexed"]
    lines = [
        "trip_index.csv is stale relative to Data/Input.",
        f"Current raw input workbooks: {freshness['input_file_count']}",
        f"Files currently present in trip_index.csv: {freshness['indexed_file_count']}",
    ]
    if not missing.empty:
        lines.append("Raw input workbook(s) not yet indexed:")
        for relative_path in missing["input_relative_path"].astype(str).tolist():
            lines.append(f"  - {relative_path}")
    lines.extend(
        [
            "Run these stages before reconciliation:",
            "  python .\\TAT_MS_INTERCITY_PIPELINE\\src\\cleaned_file_writer.py",
            "  python .\\TAT_MS_INTERCITY_PIPELINE\\src\\trip_index.py",
        ]
    )
    return "\n".join(lines)


def _rules_from_config(config):
    configured = config.get("reconciliation", {})
    defaults = {
        "maximum_log_date_difference_days": 0,
        "minimum_route_direction_score": 0.75,
        "minimum_match_score": 100.0,
        "high_confidence_assignment_margin": 10.0,
        "low_confidence_assignment_margin": 2.0,
        "maximum_boundary_sample_offset_minutes": 20.0,
        "minimum_verified_points": 2,
        "minimum_verified_distance_km": 0.05,
        "maximum_verified_duration_hours": 30.0,
        "movement_speed_threshold_kmh": 1.0,
        "movement_step_metres": 10.0,
        "boundary_stationary_extension_minutes": 15.0,
        "continuity_gap_minutes": 240.0,
        "hard_gap_minutes": 720.0,
        "maximum_implied_speed_kmh": 160.0,
        "minimum_activity_block_score": 45.0,
        "minimum_activity_score_margin": 5.0,
    }
    defaults.update(configured)
    return defaults


def _output_paths(config):
    system_folder = Path(config["project"]["results_folder"]) / "System"
    system_folder.mkdir(parents=True, exist_ok=True)
    return (
        system_folder / VERIFIED_INDEX_FILE_NAME,
        system_folder / AUDIT_FILE_NAME,
    )


def _atomic_save_csv(dataframe, output_path):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile(
        prefix=".verified_trip_index_",
        suffix=".csv",
        dir=output_path.parent,
        delete=False,
    ) as temporary_file:
        temporary_path = Path(temporary_file.name)
    try:
        dataframe.to_csv(temporary_path, index=False, encoding="utf-8-sig")
        os.replace(temporary_path, output_path)
    except Exception:
        temporary_path.unlink(missing_ok=True)
        raise


def _summary_table(
    indexed_trips,
    log_records,
    selected_matches,
    verified_records,
    unmatched_indexed,
    unmatched_log,
):
    ready = int(
        (
            verified_records.get("reconciliation_status", pd.Series(dtype=str))
            == "READY_FOR_PAIRING"
        ).sum()
    )
    data_check = len(verified_records) - ready
    duplicate_assignments = (
        int(selected_matches["source_trip_id"].duplicated().sum())
        if not selected_matches.empty
        else 0
    )
    return pd.DataFrame(
        [
            {"Metric": "Generated at", "Value": datetime.now().isoformat(timespec="seconds")},
            {"Metric": "Indexed telemetry files", "Value": len(indexed_trips)},
            {"Metric": "Trip-log rows", "Value": len(log_records)},
            {"Metric": "One-to-one log/file matches", "Value": len(selected_matches)},
            {"Metric": "Unique source files matched", "Value": selected_matches["source_trip_id"].nunique() if not selected_matches.empty else 0},
            {"Metric": "Source files assigned more than once", "Value": duplicate_assignments},
            {"Metric": "Verified trips ready for pairing", "Value": ready},
            {"Metric": "Matched trips requiring data check", "Value": data_check},
            {"Metric": "Indexed files without a log row", "Value": len(unmatched_indexed)},
            {"Metric": "Trip-log rows without a unique file", "Value": len(unmatched_log)},
            {
                "Metric": "Interpretation",
                "Value": (
                    "Every accepted source file is assigned to exactly one log row. "
                    "Time differences are resolved from full timestamps and movement; "
                    "strict/non-strict classification occurs later."
                ),
            },
        ]
    )


def _rules_table(rules, trip_log_path):
    return pd.DataFrame(
        [
            {"Rule": "Trip-log workbook", "Value": str(trip_log_path), "Purpose": "Authoritative intended-trip reference"},
            {"Rule": "Trip-log sheet", "Value": TRIP_LOG_SHEET_NAME, "Purpose": "Detailed trip records"},
            {"Rule": "Source-file assignment", "Value": "Exactly one log row per file", "Purpose": "Prevent source reuse or file splitting"},
            {"Rule": "Maximum collection-date difference", "Value": rules["maximum_log_date_difference_days"], "Purpose": "Prevent cross-date fabrication"},
            {"Rule": "Minimum route-direction score", "Value": rules["minimum_route_direction_score"], "Purpose": "Require ordered start/end route evidence"},
            {"Rule": "Telemetry time source", "Value": "Cleaned Timestamp derived from source Packet Time", "Purpose": "Actual Received Time is not used for trip boundaries or movement blocks"},
            {"Rule": "Repeated Packet Time", "Value": "Does not split a movement block", "Purpose": "Duplicate/near-duplicate packets must not create artificial trip boundaries"},
            {"Rule": "Time use during matching", "Value": "Soft evidence", "Purpose": "Extra device recording cannot exclude the correct file"},
            {"Rule": "Manual clock interpretation", "Value": "Written time plus 12-hour counterpart", "Purpose": "Resolve AM/PM transcription offsets only when Packet-Time telemetry supports the alternative"},
            {"Rule": "Exact boundary sample offset (minutes)", "Value": rules["maximum_boundary_sample_offset_minutes"], "Purpose": "Use log-anchored cut when Packet-Time samples support it"},
            {"Rule": "Fallback boundary", "Value": "Best unique movement block", "Purpose": "Resolve AM/PM, clock shifts, extra points, or missing end"},
            {"Rule": "Low speed inside verified boundary", "Value": "Preserved", "Purpose": "Retain stop, idle, congestion, and inching"},
            {"Rule": "Original/cleaned telemetry", "Value": "Never overwritten", "Purpose": "Preserve research-data lineage"},
            {"Rule": "Strict/non-strict EV-ICE label", "Value": "Not assigned here", "Purpose": "Reserved for route-overlap pairing stage"},
        ]
    )


def _format_audit_workbook(writer, tables):
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet_name, dataframe in tables.items():
        worksheet = writer.sheets[sheet_name]
        worksheet.sheet_view.showGridLines = False
        worksheet.freeze_panes = "A2"
        if worksheet.max_row >= 1 and worksheet.max_column >= 1:
            worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        worksheet.row_dimensions[1].height = 32
        for number, column_name in enumerate(dataframe.columns, start=1):
            values = [column_name] + dataframe[column_name].head(200).tolist()
            maximum = max(len(str(value)) if value is not None else 0 for value in values)
            worksheet.column_dimensions[get_column_letter(number)].width = min(
                max(maximum + 2, 12), 42
            )
            for column_cells in worksheet.iter_cols(
                min_col=number,
                max_col=number,
                min_row=2,
                max_row=max(2, worksheet.max_row),
            ):
                for cell in column_cells:
                    cell.alignment = Alignment(
                        vertical="top", wrap_text=(maximum > 40)
                    )


def _atomic_save_audit_workbook(output_path, tables):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile(
        prefix=".trip_log_reconciliation_",
        suffix=".xlsx",
        dir=output_path.parent,
        delete=False,
    ) as temporary_file:
        temporary_path = Path(temporary_file.name)
    try:
        with pd.ExcelWriter(
            temporary_path,
            engine="openpyxl",
            datetime_format="yyyy-mm-dd hh:mm:ss",
        ) as writer:
            for sheet_name, dataframe in tables.items():
                dataframe.to_excel(writer, sheet_name=sheet_name, index=False)
            _format_audit_workbook(writer, tables)
        os.replace(temporary_path, output_path)
    except Exception:
        temporary_path.unlink(missing_ok=True)
        raise


def reconcile_trip_log(config, trip_log_path=None):
    """Run the complete one-file/one-log reconciliation stage."""

    trip_log_path = Path(trip_log_path or _trip_log_path_from_config(config))
    rules = _rules_from_config(config)
    results_root = Path(config["project"]["results_folder"])
    index_path = results_root / "System" / "trip_index.csv"
    verified_index_path, audit_path = _output_paths(config)

    indexed_trips = load_file_trip_index(index_path)

    freshness = _check_trip_index_freshness(config, indexed_trips)
    if not freshness["raw_not_indexed"].empty:
        raise RuntimeError(_format_stale_index_error(freshness))

    log_records = load_trip_log(trip_log_path)
    selected_matches, unmatched_log = match_log_rows_to_indexed_files(
        indexed_trips, log_records, rules
    )
    verified_records = build_verified_trip_records(
        indexed_trips, log_records, selected_matches, rules
    )
    matched_ids = (
        set(selected_matches["source_trip_id"].astype(str))
        if not selected_matches.empty
        else set()
    )
    unmatched_indexed = indexed_trips.loc[
        ~indexed_trips["trip_id"].astype(str).isin(matched_ids)
    ].copy()
    if not unmatched_indexed.empty:
        unmatched_indexed["unmatched_reason"] = "NO_UNIQUE_TRIP_LOG_ROW_ASSIGNED"

    ready = pd.DataFrame()
    data_check = pd.DataFrame()
    if not verified_records.empty:
        ready = verified_records.loc[
            verified_records["reconciliation_status"] == "READY_FOR_PAIRING"
        ].copy()
        data_check = verified_records.loc[
            verified_records["reconciliation_status"] != "READY_FOR_PAIRING"
        ].copy()

    match_map = selected_matches.copy()
    if not match_map.empty:
        file_lookup = indexed_trips.set_index("trip_id")[
            ["file_name", "relative_path", "vehicle_type"]
        ]
        log_lookup = log_records.set_index("log_record_id")[
            [
                "log_excel_row",
                "log_date",
                "log_start_location",
                "log_end_location",
                "log_start_time_original",
                "log_end_time_original",
            ]
        ]
        match_map = match_map.join(
            file_lookup, on="source_trip_id", rsuffix="_file"
        ).join(log_lookup, on="log_record_id")

    summary = _summary_table(
        indexed_trips,
        log_records,
        selected_matches,
        verified_records,
        unmatched_indexed,
        unmatched_log,
    )
    audit_tables = {
        "Run Summary": summary,
        "One-to-One Match Map": match_map,
        "Ready Verified Trips": ready,
        "Data Check Required": data_check,
        "Unmatched Indexed Files": unmatched_indexed,
        "Unmatched Log Rows": unmatched_log,
        "Rules Applied": _rules_table(rules, trip_log_path),
    }
    _atomic_save_csv(verified_records, verified_index_path)
    _atomic_save_audit_workbook(audit_path, audit_tables)

    return {
        "indexed_files": len(indexed_trips),
        "trip_log_rows": len(log_records),
        "log_file_matches": len(selected_matches),
        "ready_verified_trips": len(ready),
        "data_check_required": len(data_check),
        "unmatched_indexed_files": len(unmatched_indexed),
        "unmatched_log_rows": len(unmatched_log),
        "duplicate_source_assignments": int(
            selected_matches["source_trip_id"].duplicated().sum()
        )
        if not selected_matches.empty
        else 0,
        "verified_index_path": str(verified_index_path),
        "audit_path": str(audit_path),
        "trip_log_path": str(trip_log_path),
        "data_check_details": (
            data_check[
                [
                    "source_file_name",
                    "log_excel_row",
                    "log_date",
                    "log_start_location",
                    "log_end_location",
                    "resolved_log_start",
                    "resolved_log_end",
                    "raw_first_timestamp",
                    "raw_last_timestamp",
                    "start_sample_offset_seconds",
                    "end_sample_offset_seconds",
                    "boundary_method",
                    "boundary_confidence",
                    "reconciliation_reason",
                ]
            ].to_dict(orient="records")
            if not data_check.empty
            else []
        ),
        "unmatched_log_details": (
            unmatched_log[
                [
                    column
                    for column in [
                        "log_excel_row",
                        "log_date",
                        "log_vehicle_type",
                        "log_start_location",
                        "log_end_location",
                        "log_start_time_original",
                        "log_end_time_original",
                        "unmatched_reason",
                        "best_match_score",
                        "best_source_trip_id",
                    ]
                    if column in unmatched_log.columns
                ]
            ].to_dict(orient="records")
            if not unmatched_log.empty
            else []
        ),
    }


if __name__ == "__main__":
    try:
        settings = load_config()
        print("Trip-log reconciliation started")
        print(f"Script: {Path(__file__).resolve()}")
        print(f"Version: {SCRIPT_VERSION}")
        result = reconcile_trip_log(settings)
        print("\nTrip-log reconciliation completed")
        print(f"Indexed telemetry files: {result['indexed_files']}")
        print(f"Trip-log rows: {result['trip_log_rows']}")
        print(f"One-to-one log/file matches: {result['log_file_matches']}")
        print(f"Verified trips ready for pairing: {result['ready_verified_trips']}")
        print(f"Matched trips requiring data check: {result['data_check_required']}")
        print(f"Indexed files without a log row: {result['unmatched_indexed_files']}")
        print(f"Trip-log rows without a unique file: {result['unmatched_log_rows']}")
        print(f"Source files assigned more than once: {result['duplicate_source_assignments']}")
        print(f"Verified trip index: {result['verified_index_path']}")
        print(f"Reconciliation audit: {result['audit_path']}")

        if result["data_check_details"]:
            print("\nDATA CHECK REQUIRED - EXACT TRIP(S)")
            for item in result["data_check_details"]:
                print(
                    "  "
                    f"Log row {item.get('log_excel_row')} | "
                    f"{item.get('log_date')} | "
                    f"{item.get('log_start_location')} -> {item.get('log_end_location')} | "
                    f"{item.get('source_file_name')}"
                )
                print(
                    "    "
                    f"Resolved log window: {item.get('resolved_log_start')} -> "
                    f"{item.get('resolved_log_end')}"
                )
                print(
                    "    "
                    f"Telemetry window: {item.get('raw_first_timestamp')} -> "
                    f"{item.get('raw_last_timestamp')}"
                )
                print(
                    "    "
                    f"Boundary offsets (s): start={item.get('start_sample_offset_seconds')} | "
                    f"end={item.get('end_sample_offset_seconds')}"
                )
                print(
                    "    "
                    f"Method/confidence: {item.get('boundary_method')} / "
                    f"{item.get('boundary_confidence')}"
                )
                print(f"    Reason: {item.get('reconciliation_reason')}")

        if result["unmatched_log_details"]:
            print("\nTRIP-LOG ROWS WITHOUT A UNIQUE FILE")
            for item in result["unmatched_log_details"]:
                print(
                    "  "
                    f"Log row {item.get('log_excel_row')} | "
                    f"{item.get('log_date')} | "
                    f"{item.get('log_vehicle_type')} | "
                    f"{item.get('log_start_location')} -> {item.get('log_end_location')} | "
                    f"{item.get('unmatched_reason')}"
                )
    except Exception as error:
        print(f"\nTRIP-LOG RECONCILIATION FAILED | {error}")
        raise SystemExit(1)