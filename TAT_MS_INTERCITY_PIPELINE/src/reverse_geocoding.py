"""Add Reverse Geocoded Address to paired EV/ICE sheets.

Run after feature_engineering.py and before phase_segmentation.py:

    python src/reverse_geocoding.py

Only stationary rows are reverse geocoded. Coordinates are rounded and cached
in Results/System/reverse_geocode_cache.csv, so repeated locations do not
create repeated API requests.
"""

from __future__ import annotations

import json
import os
import stat
import time
import urllib.parse
import urllib.request
from datetime import datetime
from pathlib import Path
from tempfile import NamedTemporaryFile

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config_loader import load_config


SCRIPT_VERSION = "2026-08-06_REVERSE_GEOCODING_V1"
PAIR_INDEX_NAME = "pair_index.csv"
CACHE_NAME = "reverse_geocode_cache.csv"
EV_SHEET = "01_EV_Cropped_Data"
ICE_SHEET = "02_ICE_Cropped_Data"
LAT = "Latitude"
LON = "Longitude"
TIME_GAP = "Time Difference (s)"
DISTANCE = "Haversine Distance (m)"
SPEED = "Speed (km/h)"
OUTPUT_COLUMN = "Reverse Geocoded Address"
HEADER_BLUE = "1F4E78"


def clean_text(value) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def read_csv(path: Path) -> pd.DataFrame:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return pd.read_csv(path, encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"Could not decode {path}")


def rules(config: dict) -> dict:
    phase = config.get("phase_segmentation", {}) or {}
    return {
        "decimals": int(phase.get("station_cache_coordinate_decimals", 4)),
        "delay": float(phase.get("station_lookup_delay_seconds", 1.1)),
        "timeout": float(phase.get("station_lookup_timeout_seconds", 15.0)),
        "user_agent": clean_text(
            phase.get("nominatim_user_agent")
        ) or "tat-ms-research-pipeline/1.0",
        "idle_seconds": float(phase.get("idle_gap_min_seconds", 30.0)),
        "stationary_distance_m": float(
            phase.get("stationary_displacement_m", 50.0)
        ),
        "zero_speed_kmh": float(
            phase.get("zero_speed_tolerance_kmh", 0.01)
        ),
    }


class GeocodeCache:
    COLUMNS = [
        "cache_key", "latitude", "longitude", "address",
        "lookup_status", "updated_at",
    ]

    def __init__(self, path: Path, settings: dict):
        self.path = path
        self.settings = settings
        self.addresses = {}
        self.rows = {}
        self.last_request = 0.0
        self.dirty = 0
        self.hits = 0
        self.requests = 0
        if path.exists():
            frame = read_csv(path)
            for _, row in frame.iterrows():
                key = clean_text(row.get("cache_key"))
                if key:
                    self.addresses[key] = clean_text(row.get("address"))
                    self.rows[key] = row.to_dict()

    def key(self, latitude: float, longitude: float) -> str:
        d = self.settings["decimals"]
        return f"{latitude:.{d}f}|{longitude:.{d}f}"

    def resolve(self, latitude: float, longitude: float) -> str:
        key = self.key(latitude, longitude)
        if key in self.addresses:
            self.hits += 1
            return self.addresses[key]

        elapsed = time.monotonic() - self.last_request
        delay = self.settings["delay"]
        if elapsed < delay:
            time.sleep(delay - elapsed)

        query = urllib.parse.urlencode({
            "format": "jsonv2",
            "lat": f"{latitude:.7f}",
            "lon": f"{longitude:.7f}",
            "zoom": "18",
            "addressdetails": "1",
        })
        request = urllib.request.Request(
            f"https://nominatim.openstreetmap.org/reverse?{query}",
            headers={"User-Agent": self.settings["user_agent"]},
        )
        self.requests += 1
        self.last_request = time.monotonic()

        try:
            with urllib.request.urlopen(
                request,
                timeout=self.settings["timeout"],
            ) as response:
                payload = json.loads(response.read().decode("utf-8"))
            address = clean_text(payload.get("display_name"))
        except Exception:
            address = ""

        self.addresses[key] = address
        self.rows[key] = {
            "cache_key": key,
            "latitude": latitude,
            "longitude": longitude,
            "address": address,
            "lookup_status": "MATCHED" if address else "NO_ADDRESS_RETURNED",
            "updated_at": datetime.now().isoformat(timespec="seconds"),
        }
        self.dirty += 1
        if self.dirty >= 20:
            self.save()
        return address

    def save(self) -> None:
        if not self.dirty:
            return
        self.path.parent.mkdir(parents=True, exist_ok=True)
        frame = pd.DataFrame(list(self.rows.values()), columns=self.COLUMNS)
        with NamedTemporaryFile(
            prefix=f".{self.path.stem}_",
            suffix=".csv",
            dir=self.path.parent,
            delete=False,
        ) as temporary_file:
            temporary = Path(temporary_file.name)
        try:
            frame.to_csv(temporary, index=False, encoding="utf-8-sig")
            os.replace(temporary, self.path)
            self.dirty = 0
        except Exception:
            temporary.unlink(missing_ok=True)
            raise


def workbook_path(row: pd.Series, results_root: Path) -> Path:
    configured = clean_text(row.get("pair_workbook_path"))
    if configured and Path(configured).exists():
        return Path(configured)

    pair_id = clean_text(row.get("pair_id"))
    matches = [
        path for path in (results_root / "Paired").rglob(f"{pair_id}_*.xlsx")
        if not path.name.startswith(("~$", "."))
    ]
    if not matches:
        raise FileNotFoundError(f"Pair workbook not found for {pair_id}")
    return max(matches, key=lambda path: path.stat().st_mtime)


def geocode_frame(
    frame: pd.DataFrame,
    cache: GeocodeCache,
    settings: dict,
) -> tuple[list[str], int, int, int]:
    required = [LAT, LON, TIME_GAP, DISTANCE, SPEED]
    missing = [column for column in required if column not in frame.columns]
    if missing:
        raise ValueError("Missing feature columns: " + ", ".join(missing))

    gap = pd.to_numeric(frame[TIME_GAP], errors="coerce").fillna(0.0)
    distance = pd.to_numeric(frame[DISTANCE], errors="coerce").fillna(np.inf)
    speed = pd.to_numeric(frame[SPEED], errors="coerce").fillna(np.inf)
    latitude = pd.to_numeric(frame[LAT], errors="coerce")
    longitude = pd.to_numeric(frame[LON], errors="coerce")

    stationary = (
        (gap >= settings["idle_seconds"])
        & (distance <= settings["stationary_distance_m"])
        & (speed <= settings["zero_speed_kmh"])
    )

    output = [""] * len(frame)
    before_hits = cache.hits
    before_requests = cache.requests

    for position in np.flatnonzero(stationary.to_numpy()):
        lat = latitude.iloc[position]
        lon = longitude.iloc[position]
        if pd.notna(lat) and pd.notna(lon):
            output[position] = cache.resolve(float(lat), float(lon))

    return (
        output,
        int(stationary.sum()),
        cache.hits - before_hits,
        cache.requests - before_requests,
    )


def write_column(path: Path, ev_values: list[str], ice_values: list[str]) -> None:
    workbook = load_workbook(path)
    temporary = None
    try:
        for sheet_name, values in ((EV_SHEET, ev_values), (ICE_SHEET, ice_values)):
            worksheet = workbook[sheet_name]
            headers = {
                clean_text(cell.value): cell.column
                for cell in worksheet[1]
                if clean_text(cell.value)
            }
            column = headers.get(OUTPUT_COLUMN, worksheet.max_column + 1)
            worksheet.cell(1, column, OUTPUT_COLUMN)
            cell = worksheet.cell(1, column)
            cell.fill = PatternFill("solid", fgColor=HEADER_BLUE)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )

            if worksheet.max_row - 1 != len(values):
                raise RuntimeError(f"{sheet_name} row count changed")

            for row_number, value in enumerate(values, start=2):
                worksheet.cell(row_number, column, value)

            worksheet.column_dimensions[
                get_column_letter(column)
            ].width = 55
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions

        with NamedTemporaryFile(
            prefix=f".{path.stem}_geocode_",
            suffix=".xlsx",
            dir=path.parent,
            delete=False,
        ) as file:
            temporary = Path(file.name)

        workbook.save(temporary)
        workbook.close()
        workbook = None

        try:
            os.chmod(path, stat.S_IREAD | stat.S_IWRITE)
        except OSError:
            pass
        os.replace(temporary, path)
    finally:
        if workbook is not None:
            workbook.close()
        if temporary is not None and temporary.exists():
            temporary.unlink(missing_ok=True)


def duration(seconds: float) -> str:
    seconds = max(0, int(round(seconds)))
    hours, remainder = divmod(seconds, 3600)
    minutes, seconds = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{seconds:02d}"


def run(config: dict) -> None:
    results_root = Path(config["project"]["results_folder"])
    pair_index = read_csv(results_root / "System" / PAIR_INDEX_NAME)
    settings = rules(config)
    cache = GeocodeCache(
        results_root / "System" / CACHE_NAME,
        settings,
    )

    started = time.monotonic()
    print("=" * 76)
    print("REVERSE GEOCODING")
    print("=" * 76)
    print(f"Pair workbooks: {len(pair_index)}")
    print(f"Cache entries loaded: {len(cache.addresses)}")
    print()

    for sequence, (_, row) in enumerate(pair_index.iterrows(), start=1):
        pair_started = time.monotonic()
        path = workbook_path(row, results_root)
        pair_id = clean_text(row.get("pair_id"))
        print(f"[{sequence}/{len(pair_index)}] {pair_id} | {path.name}", flush=True)

        with pd.ExcelFile(path, engine="openpyxl") as excel:
            ev = pd.read_excel(excel, sheet_name=EV_SHEET)
            ice = pd.read_excel(excel, sheet_name=ICE_SHEET)

        ev_values, ev_stationary, ev_hits, ev_requests = geocode_frame(
            ev, cache, settings
        )
        ice_values, ice_stationary, ice_hits, ice_requests = geocode_frame(
            ice, cache, settings
        )
        write_column(path, ev_values, ice_values)

        elapsed = time.monotonic() - started
        eta = elapsed / sequence * (len(pair_index) - sequence)
        print(
            f"    Stationary rows | EV: {ev_stationary} | ICE: {ice_stationary}",
            flush=True,
        )
        print(
            f"    Cache hits: {ev_hits + ice_hits} | "
            f"New requests: {ev_requests + ice_requests}",
            flush=True,
        )
        print(
            f"    This pair: {duration(time.monotonic() - pair_started)} | "
            f"Elapsed: {duration(elapsed)} | ETA: {duration(eta)}",
            flush=True,
        )
        print("-" * 76, flush=True)

    cache.save()
    print()
    print("REVERSE GEOCODING COMPLETED")
    print(f"Cache hits: {cache.hits}")
    print(f"API requests: {cache.requests}")
    print(f"Cache: {cache.path}")


if __name__ == "__main__":
    try:
        run(load_config())
    except Exception as error:
        print(f"\nREVERSE GEOCODING FAILED | {error}")
        raise SystemExit(1)