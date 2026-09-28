import os
import shutil
from datetime import datetime
from pathlib import Path
from tempfile import NamedTemporaryFile

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config_loader import load_config
from data_cleaner import (
    clean_telemetry_data_with_audit,
    write_cleaning_audit_report,
)
from file_discovery import discover_input_files
from input_validator import validate_input_file


CLEANED_SHEET_NAME = "Cleaned Data"


def _normalise_relative_path(path_value):
    """Return a stable, case-insensitive relative-path key."""

    return str(path_value).replace("\\", "/").rstrip("/").lower()


def build_cleaned_output_path(item, config):
    """Return the cleaned path while preserving the input subfolders."""

    input_folder = Path(config["project"]["input_folder"])
    results_folder = Path(config["project"]["results_folder"])
    source_path = Path(item["file_path"])

    try:
        relative_path = source_path.resolve().relative_to(
            input_folder.resolve()
        )
    except ValueError as error:
        raise ValueError(
            f"Input file is outside the configured input folder: "
            f"{source_path}"
        ) from error

    return results_folder / "Cleaned" / relative_path


def _format_cleaned_worksheet(worksheet):
    """Apply readable formatting without changing any data values."""

    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )
    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    worksheet.row_dimensions[1].height = 32

    column_widths = {
        "Timestamp": 21,
        "Latitude": 14,
        "Longitude": 14,
        "Speed(m/s)": 14,
        "Speed(km/h)": 15,
        "Source Excel Row": 18,
        "Original Speed(m/s)": 21,
        "Original Speed(km/h)": 22,
        "Speed Recovery Status": 24,
        "Address": 55,
    }

    header_by_column = {
        cell.column: cell.value
        for cell in worksheet[1]
    }

    for column_number in range(1, worksheet.max_column + 1):
        header = header_by_column.get(column_number)
        column_letter = get_column_letter(column_number)
        worksheet.column_dimensions[column_letter].width = (
            column_widths.get(header, 16)
        )

        for cell in worksheet.iter_cols(
            min_col=column_number,
            max_col=column_number,
            min_row=2,
        ):
            for value_cell in cell:
                value_cell.alignment = Alignment(
                    vertical="top",
                )

                if header == "Timestamp":
                    value_cell.number_format = (
                        "yyyy-mm-dd hh:mm:ss"
                    )
                elif header in {"Latitude", "Longitude"}:
                    value_cell.number_format = "0.000000"
                elif header in {
                    "Speed(m/s)",
                    "Speed(km/h)",
                    "Original Speed(m/s)",
                    "Original Speed(km/h)",
                }:
                    value_cell.number_format = "0.000"
                elif header == "Source Excel Row":
                    value_cell.number_format = "0"
                elif header == "Address":
                    value_cell.alignment = Alignment(
                        vertical="top",
                        wrap_text=True,
                    )


def save_cleaned_workbook(dataframe, output_path):
    """
    Save one cleaned workbook safely.

    The workbook is first written to a temporary file. It replaces an older
    cleaned copy only after the new workbook has been written successfully.
    The raw input workbook is never changed.
    """

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with NamedTemporaryFile(
        prefix=f".{output_path.stem}_",
        suffix=".xlsx",
        dir=output_path.parent,
        delete=False,
    ) as temporary_file:
        temporary_path = Path(temporary_file.name)

    try:
        with pd.ExcelWriter(
            temporary_path,
            engine="openpyxl",
            datetime_format="yyyy-mm-dd hh:mm:ss",
        ) as writer:
            dataframe.to_excel(
                writer,
                sheet_name=CLEANED_SHEET_NAME,
                index=False,
            )

            _format_cleaned_worksheet(
                writer.book[CLEANED_SHEET_NAME]
            )

        os.replace(temporary_path, output_path)

    except Exception:
        temporary_path.unlink(missing_ok=True)
        raise


def _validation_failure_summary(item, validation):
    """Create an audit summary for a validator-rejected input file."""

    return {
        "file_name": item["file_name"],
        "source_file": str(item["file_path"]),
        "source_schema": validation.get("source_schema"),
        "source_sheet": None,
        "processing_status": "SKIPPED_VALIDATION",
        "original_rows": validation.get("row_count", 0),
        "invalid_speeds_recovered": 0,
        "invalid_speeds_unresolved": validation.get(
            "invalid_speeds",
            0,
        ),
        "unrecoverable_rows_removed": 0,
        "duplicate_rows_removed": 0,
        "total_rows_removed": 0,
        "cleaned_rows": 0,
        "notes": "; ".join(validation.get("issues", [])),
    }


def _failed_file_summary(item, error):
    """Create an audit summary for an input that failed unexpectedly."""

    return {
        "file_name": item["file_name"],
        "source_file": str(item["file_path"]),
        "source_schema": None,
        "source_sheet": None,
        "processing_status": "FAILED",
        "original_rows": 0,
        "invalid_speeds_recovered": 0,
        "invalid_speeds_unresolved": 0,
        "unrecoverable_rows_removed": 0,
        "duplicate_rows_removed": 0,
        "total_rows_removed": 0,
        "cleaned_rows": 0,
        "notes": str(error),
    }


def _remove_empty_directories(root_folder):
    """Remove empty subfolders while preserving the cleaned root."""

    root_folder = Path(root_folder)

    if not root_folder.exists():
        return

    directories = sorted(
        (path for path in root_folder.rglob("*") if path.is_dir()),
        key=lambda path: len(path.parts),
        reverse=True,
    )

    for directory in directories:
        try:
            directory.rmdir()
        except OSError:
            # The directory is not empty, so it must be retained.
            pass


def archive_obsolete_cleaned_files(
    config,
    current_relative_paths,
):
    """
    Move cleaned files that no longer correspond to a valid current input.

    The files are moved outside Results/Cleaned instead of being deleted, so
    an accidental source deletion or rename remains recoverable.
    """

    results_root = Path(config["project"]["results_folder"])
    cleaned_root = results_root / "Cleaned"

    if not cleaned_root.exists():
        return {
            "archived_count": 0,
            "failed_count": 0,
            "archive_root": None,
        }

    accepted_extensions = {
        extension.lower()
        for extension in config["files"]["accepted_extensions"]
    }
    current_keys = {
        _normalise_relative_path(path)
        for path in current_relative_paths
    }

    obsolete_files = []

    for path in cleaned_root.rglob("*"):
        if not path.is_file():
            continue

        if path.name.startswith(("~$", ".")):
            continue

        if path.suffix.lower() not in accepted_extensions:
            continue

        relative_path = path.relative_to(cleaned_root)

        if _normalise_relative_path(relative_path) not in current_keys:
            obsolete_files.append((path, relative_path))

    if not obsolete_files:
        return {
            "archived_count": 0,
            "failed_count": 0,
            "archive_root": None,
        }

    run_timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    archive_root = (
        results_root
        / "System"
        / "Archive"
        / "Cleaned"
        / run_timestamp
    )

    archived_count = 0
    failed_count = 0

    for source_path, relative_path in obsolete_files:
        destination_path = archive_root / relative_path

        try:
            destination_path.parent.mkdir(
                parents=True,
                exist_ok=True,
            )
            shutil.move(str(source_path), str(destination_path))
            archived_count += 1
            print(
                f"ARCHIVED OBSOLETE | {relative_path} | "
                f"Backup: {destination_path}"
            )
        except Exception as error:
            failed_count += 1
            print(
                f"ARCHIVE FAILED | {relative_path} | {error}"
            )

    _remove_empty_directories(cleaned_root)

    return {
        "archived_count": archived_count,
        "failed_count": failed_count,
        "archive_root": (
            str(archive_root) if archived_count > 0 else None
        ),
    }


def _has_managed_cleaned_files(config):
    """Return whether the cleaned folder contains telemetry output files."""

    cleaned_root = (
        Path(config["project"]["results_folder"])
        / "Cleaned"
    )

    if not cleaned_root.exists():
        return False

    accepted_extensions = {
        extension.lower()
        for extension in config["files"]["accepted_extensions"]
    }

    return any(
        path.is_file()
        and not path.name.startswith(("~$", "."))
        and path.suffix.lower() in accepted_extensions
        for path in cleaned_root.rglob("*")
    )


def write_all_cleaned_files(config):
    """Clean inputs, rebuild the audit, and synchronize cleaned outputs."""

    discovered_files = discover_input_files(config)

    print("Cleaned-file saving started")
    print(f"Files discovered: {len(discovered_files)}")

    saved_count = 0
    skipped_count = 0
    failed_count = 0
    total_rows_saved = 0
    current_relative_paths = set()
    file_summaries = []
    removed_rows_tables = []
    speed_audit_tables = []

    for item in discovered_files:
        try:
            validation = validate_input_file(
                item["file_path"],
                config,
            )

            if validation["status"] != "VALID":
                skipped_count += 1
                file_summaries.append(
                    _validation_failure_summary(item, validation)
                )

                print(
                    f"SKIPPED | {item['file_name']} | "
                    "Input validation failed"
                )
                continue

            (
                cleaned_data,
                summary,
                removed_rows,
                speed_audit,
            ) = clean_telemetry_data_with_audit(
                item["file_path"], config
            )

            minimum_valid_rows = config["cleaning"][
                "minimum_valid_rows"
            ]

            if len(cleaned_data) < minimum_valid_rows:
                skipped_count += 1
                summary["processing_status"] = "SKIPPED_TOO_FEW_ROWS"
                summary["notes"] = (
                    "Insufficient valid rows after cleaning: "
                    f"{len(cleaned_data)} < {minimum_valid_rows}"
                )
                file_summaries.append(summary)
                removed_rows_tables.append(removed_rows)
                speed_audit_tables.append(speed_audit)

                print(
                    f"SKIPPED | {item['file_name']} | "
                    "Insufficient valid rows after cleaning"
                )
                continue

            output_path = build_cleaned_output_path(
                item,
                config,
            )

            save_cleaned_workbook(
                cleaned_data,
                output_path,
            )

            saved_count += 1
            total_rows_saved += len(cleaned_data)
            current_relative_paths.add(
                output_path.relative_to(
                    Path(config["project"]["results_folder"])
                    / "Cleaned"
                )
            )
            file_summaries.append(summary)
            removed_rows_tables.append(removed_rows)
            speed_audit_tables.append(speed_audit)

            print(
                f"SAVED | {item['file_name']} | "
                f"Rows: {len(cleaned_data)} | "
                f"Recovered speeds: "
                f"{summary['invalid_speeds_recovered']} | "
                f"Unresolved speeds: "
                f"{summary['invalid_speeds_unresolved']} | "
                f"Output: {output_path}"
            )

        except PermissionError as error:
            failed_count += 1
            file_summaries.append(
                _failed_file_summary(item, error)
            )

            print(
                f"FAILED | {item['file_name']} | "
                "Close its existing cleaned workbook in Excel "
                "and run this program again."
            )

        except Exception as error:
            failed_count += 1
            file_summaries.append(
                _failed_file_summary(item, error)
            )
            print(f"FAILED | {item['file_name']} | {error}")

    cleaned_root = (
        Path(config["project"]["results_folder"])
        / "Cleaned"
    )

    archive_summary = {
        "archived_count": 0,
        "failed_count": 0,
        "archive_root": None,
    }
    safety_guard_triggered = False

    # Synchronize only after every input was processed without an unexpected
    # failure. Validation-skipped files are intentionally excluded, so an old
    # cleaned copy cannot remain active after its source becomes invalid.
    if (
        len(discovered_files) == 0
        and _has_managed_cleaned_files(config)
    ):
        safety_guard_triggered = True
        print(
            "CLEANED FOLDER NOT SYNCHRONIZED | No input files were "
            "discovered, but existing cleaned outputs were found. This "
            "safety check prevents a mistaken empty input path from "
            "archiving every cleaned trip."
        )
    elif failed_count == 0:
        archive_summary = archive_obsolete_cleaned_files(
            config,
            current_relative_paths,
        )
    else:
        print(
            "CLEANED FOLDER NOT SYNCHRONIZED | One or more inputs "
            "failed; existing outputs were retained for safety."
        )

    audit_saved = False
    audit_report_path = None

    try:
        audit_report_path = write_cleaning_audit_report(
            config=config,
            file_summaries=file_summaries,
            removed_rows=removed_rows_tables,
            speed_audit_rows=speed_audit_tables,
            files_discovered=len(discovered_files),
            cleaned_count=saved_count,
            skipped_count=skipped_count,
            failed_count=(
                failed_count + archive_summary["failed_count"]
            ),
        )
        audit_saved = True
        print(f"Cleaning audit saved: {audit_report_path}")
    except PermissionError:
        print(
            "AUDIT REPORT NOT SAVED | Close "
            "Data_Cleaning_Audit.xlsx if it is open in Excel, "
            "then run the updater again."
        )
    except Exception as error:
        print(f"AUDIT REPORT FAILED | {error}")

    sync_succeeded = (
        failed_count == 0
        and archive_summary["failed_count"] == 0
        and audit_saved
        and not safety_guard_triggered
    )

    print(f"Files saved successfully: {saved_count}")
    print(f"Files skipped: {skipped_count}")
    print(f"Files failed: {failed_count}")
    print(
        "Obsolete cleaned files archived: "
        f"{archive_summary['archived_count']}"
    )
    print(f"Total cleaned rows saved: {total_rows_saved}")
    print(f"Cleaned files folder: {cleaned_root}")

    return {
        "files_discovered": len(discovered_files),
        "files_saved": saved_count,
        "files_skipped": skipped_count,
        "files_failed": failed_count,
        "archive_failures": archive_summary["failed_count"],
        "obsolete_files_archived": archive_summary[
            "archived_count"
        ],
        "archive_root": archive_summary["archive_root"],
        "audit_saved": audit_saved,
        "audit_report_path": (
            str(audit_report_path)
            if audit_report_path is not None
            else None
        ),
        "sync_succeeded": sync_succeeded,
        "safety_guard_triggered": safety_guard_triggered,
        "total_rows_saved": total_rows_saved,
        "cleaned_root": str(cleaned_root),
    }


if __name__ == "__main__":
    settings = load_config()
    write_all_cleaned_files(settings)